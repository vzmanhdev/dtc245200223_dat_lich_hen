"""Report that follows docx/Bao_Cao_Trien_Khai_He_Thong.docx section by section.
Keeps the approved institutional cover and the deployment template body formatting; content comes from this repo.
Requires python-docx and Pillow. Afterwards open the output in Word and update the table of
contents (References > Update Table) so it shows real page numbers.
"""
from pathlib import Path
import subprocess, json
from copy import deepcopy
from io import BytesIO
from hashlib import sha256
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'docx/Bao_Cao_Trien_Khai_He_Thong.docx'
OUT = ROOT / 'reports/Bao_cao_Trien_khai_Theo_mau_dtc245200223.docx'
E = ROOT / 'docs/evidence'
D = E / 'diagrams'
NAVY, BLUE, BLACK = RGBColor(0x00, 0x33, 0x66), RGBColor(0x33, 0x66, 0x99), RGBColor(0, 0, 0)
TNR = 'Times New Roman'

d = Document(TEMPLATE)
body = d.element.body
paras = d.paragraphs


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ''


def fmt(r, size=13, bold=False, italic=False, color=BLACK, font=TNR):
    r.font.name = font
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


# Cover: keep layout, fill placeholders.
set_text(paras[3], 'Đề tài: ')
fmt(paras[3].add_run('Đề 40 – Hệ thống Đặt lịch hẹn (Appointment)'), 16, True, color=NAVY)
topic = paras[4]
fmt(topic.add_run('Website đặt lịch hẹn: dịch vụ, khung giờ, khách hàng, xác nhận – '
                  'triển khai bằng Docker Compose với PostgreSQL, pgAdmin, Nginx, '
                  'Prometheus, Grafana, Loki và Promtail'), 13, italic=True)
topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover = paras[5]
lines = ['Sinh viên thực hiện: Nguyễn Văn Mạnh', 'Mã số sinh viên (MSSV): dtc245200223',
         'Lớp / Khóa: CNTTK23B', 'Repository: github.com/vzmanhdev/dtc245200223_dat_lich_hen']
for r in cover.runs:
    r.text = ''
for i, line in enumerate(lines):
    r = fmt(cover.add_run(line), 13, italic=True)
    if i < len(lines) - 1:
        r.add_break()
set_text(paras[7], 'Thái Nguyên, ngày 25 tháng 10 năm 2026')

# Drop the template body after the cover page break; the section properties stay.
cover_end = paras[8]._p
started = False
for el in list(body):
    if started and el.tag != qn('w:sectPr'):
        body.remove(el)
    if el is cover_end:
        started = True

sec = d.sections[0]
sec.different_first_page_header_footer = True


def field(p, code):
    for kind, text in (('begin', None), (None, code), ('separate', None), (None, '1'), ('end', None)):
        r = p.add_run()
        if kind:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
            r._element.append(el)
        elif text == code:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = code
            r._element.append(el)
        else:
            r.text = text
        fmt(r, 11)


foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(foot, 'PAGE')


def outline(p, level):
    el = OxmlElement('w:outlineLvl')
    el.set(qn('w:val'), str(level))
    p._p.get_or_add_pPr().append(el)


def para(text='', size=13, bold=False, italic=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=120, keep=False):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before / 20), Pt(after / 20), 1.15
    pf.keep_with_next = keep
    p.alignment = align
    if text:
        rich(p, text, size, bold, italic, color)
    return p


def rich(p, text, size=13, bold=False, italic=False, color=BLACK):
    # **x** = bold, `x` = monospace fragments inside body text.
    import re
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**'):
            fmt(p.add_run(part[2:-2]), size, True, italic, color)
        elif part.startswith('`'):
            fmt(p.add_run(part[1:-1]), size - 1.5, bold, italic, color, 'Consolas')
        else:
            fmt(p.add_run(part), size, bold, italic, color)


def chapter(text):
    p = para(text, 14, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, before=360, after=120, keep=True)
    p.paragraph_format.page_break_before = True
    outline(p, 0)


def section(text):
    p = para(text, 13, True, True, BLUE, WD_ALIGN_PARAGRAPH.LEFT, 240, 80, True)
    outline(p, 1)


def sub(text):
    p = para(text, 13, True, align=WD_ALIGN_PARAGRAPH.LEFT, before=160, after=60, keep=True)
    outline(p, 2)


def bullet(label, text=''):
    p = d.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        rich(p, f'**{label}:** {text}')
    else:
        rich(p, label)


def code(*lines):
    for i, line in enumerate(lines):
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(6 if i == 0 else 0)
        pf.space_after = Pt(6 if i == len(lines) - 1 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        pf.keep_with_next = i < len(lines) - 1
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F2F2F2')
        # Schema order: shd must precede spacing/ind or Word asks to repair the file.
        p._p.get_or_add_pPr().insert_element_before(shd, 'w:tabs', 'w:spacing', 'w:ind', 'w:jc', 'w:rPr')
        fmt(p.add_run(line or ' '), 9.5, font='Consolas')


fig_no = {}
figure_titles=[]
table_titles=[]


def figure(path, caption, width=15.0, max_height=19.0):
    w, h = Image.open(path).size
    width = min(width, max_height * w / h)
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=120, after=40, keep=True)
    p.add_run().add_picture(str(path), width=Cm(width))
    ch = current_chapter[0]
    fig_no[ch] = fig_no.get(ch, 0) + 1
    q=para('Hình ',12,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=60)
    field(q,'SEQ Hinh \\* ARABIC');fmt(q.add_run('. '+caption),12,italic=True)
    figure_titles.append(caption)
    note='Nguồn: tác giả dựng theo source/config.' if path.parent==D else 'Nguồn: ảnh chụp hệ thống thực hành; số liệu phản ánh thời điểm chụp.'
    para(note,10.5,italic=True,align=WD_ALIGN_PARAGRAPH.LEFT,after=120)


tab_no = {}


def table(headers, rows, widths, caption):
    ch = current_chapter[0]
    tab_no[ch] = tab_no.get(ch, 0) + 1
    q=para('Bảng ',12,True,align=WD_ALIGN_PARAGRAPH.CENTER,before=120,after=60,keep=True)
    field(q,'SEQ Bang \\* ARABIC');fmt(q.add_run('. '+caption),12,True)
    table_titles.append(caption)
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, wcm in zip(t.columns, widths):
        col.width = Cm(wcm)
    for row_cells, values, head in [(t.rows[0].cells, headers, True)] + \
            [(t.add_row().cells, r, False) for r in rows]:
        for c, v, wcm in zip(row_cells, values, widths):
            c.width = Cm(wcm)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = head
            rich(p, v, 11, bold=head)
            if head:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'DCE6F1')
                c._tc.get_or_add_tcPr().append(shd)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    head_row = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    head_row.append(el)
    para(after=60)


current_chapter = [0]

# Table of contents (Word fills it when the render script updates fields).
p = para('MỤC LỤC', 14, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=200)
toc = para(align=WD_ALIGN_PARAGRAPH.LEFT)
field(toc, 'TOC \\o "1-2" \\h \\z \\u')

# Separate automatic figure and table lists, as requested.
para('DANH MỤC HÌNH ẢNH',14,True,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,after=160).paragraph_format.page_break_before=True
field(para(align=WD_ALIGN_PARAGRAPH.LEFT),'TOC \\h \\z \\c "Hinh"')
para('DANH MỤC BẢNG',14,True,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,after=160).paragraph_format.page_break_before=True
field(para(align=WD_ALIGN_PARAGRAPH.LEFT),'TOC \\h \\z \\c "Bang"')

# ---------------------------------------------------------------- Chapter 1
current_chapter[0] = 1
chapter('CHƯƠNG 1: TỔNG QUAN KIẾN TRÚC HỆ THỐNG VÀ CHUẨN BỊ MÔI TRƯỜNG')
section('1.1. Sơ đồ kiến trúc tổng thể')
para('Đề 40 yêu cầu triển khai website đặt lịch hẹn gồm các nghiệp vụ dịch vụ, khung giờ, khách hàng và '
     'xác nhận, sử dụng PostgreSQL + pgAdmin cùng bộ công cụ vận hành chung (GitHub, Nginx reverse proxy, '
     'Prometheus + Grafana, Loki + LogQL, hardening). Ứng dụng được viết bằng **Node.js 24 + Express 5**, giao diện '
     'render phía server bằng **EJS**, dữ liệu lưu trong **PostgreSQL 17**. Toàn bộ hệ thống được mô tả trong một '
     'file `compose.yaml` duy nhất và chạy trên máy cục bộ (Windows 11 + Docker Desktop/WSL 2).')
para('Hệ thống được tổ chức theo mô hình **container hóa nhiều dịch vụ**: một backend nghiệp vụ duy nhất được tách '
     'module bên trong, bao quanh bởi các dịch vụ hạ tầng độc lập (proxy, cơ sở dữ liệu, giám sát, log). Các thành '
     'phần giao tiếp với nhau qua bốn Docker network riêng, trong đó hai mạng `db` và `monitoring` được đặt '
     '`internal: true` để hạn chế kết nối ngoài ở phạm vi từng network; service gắn nhiều mạng vẫn cần kiểm soát riêng. Compose gồm **13 service**: 11 service chạy dài hạn và 2 job '
     'chạy một lần (`log-init`, `migrate`) kết thúc với mã thoát 0.')
para('Sơ đồ tổng thể bên dưới chia thành ba luồng chính, gồm các thành phần:')
bullet('Nginx Reverse Proxy (nginx)', 'cổng giao tiếp duy nhất của website. Nhận request từ trình duyệt qua '
       'HTTPS `127.0.0.1:443` (HTTP `8080` là cổng phụ), kết thúc TLS, gắn security headers, giới hạn tốc độ và '
       'chuyển tiếp vào `app:3000`. Image `nginx-unprivileged` chạy không cần quyền root.')
bullet('Ứng dụng Web & Database (app, postgres, migrate, pgadmin)', 'container `app` xử lý nghiệp vụ, phiên đăng '
       'nhập và kết nối `postgres:5432` qua mạng `db` nội bộ. Job `migrate` áp dụng schema v1/v2/v3 và tạo các role '
       'cơ sở dữ liệu trước khi app khởi động. pgAdmin là công cụ quản trị DB trực quan, chỉ mở trên '
       '`127.0.0.1:5050`. PostgreSQL không publish cổng ra máy host.')
bullet('Giám sát Metrics (Prometheus + Grafana)', 'Prometheus chủ động scrape mỗi 15 giây năm job: chính nó, '
       'endpoint `/metrics` của app, `nginx-exporter`, `postgres-exporter` và `cadvisor` (CPU/RAM/Network của '
       'từng container). Grafana đọc dữ liệu bằng PromQL và hiển thị dashboard được cấp sẵn (provisioning).')
bullet('Quản lý Log tập trung (Loki + Promtail)', 'app, Nginx và PostgreSQL ghi log ra ba volume riêng; Promtail '
       'đọc các file này ở chế độ chỉ đọc, gắn nhãn `service`/`environment`, phân tích thời gian sự kiện rồi đẩy về '
       'Loki. Quản trị viên tra cứu log trong Grafana Explore bằng LogQL.')
figure(D / '00-full-system-architecture.png', 'Sơ đồ kiến trúc tổng thể 13 service Docker Compose', 14.5, 18.5)
table(['Service', 'Image (pin digest)', 'Mạng', 'Vai trò / truy cập'], [
    ['log-init', 'busybox 1.37.0', 'none', 'Job cấp quyền thư mục log, Exited 0'],
    ['postgres', 'postgres:17-bookworm', 'db', 'Lưu dữ liệu nghiệp vụ và session; không publish'],
    ['migrate', 'manh-appointment:local', 'db', 'Job áp dụng schema/role, Exited 0'],
    ['app', 'manh-appointment:local', 'frontend, db, monitoring', 'Express/EJS, `app:3000`, `/metrics`'],
    ['nginx', 'nginx-unprivileged stable-alpine', 'frontend, monitoring', 'TLS + proxy, host 443/8080'],
    ['pgadmin', 'dpage/pgadmin4:9', 'db, management', 'Quản trị DB, host 5050'],
    ['postgres-exporter', 'postgres-exporter v0.17.1', 'db, monitoring', 'Metrics PostgreSQL :9187'],
    ['nginx-exporter', 'nginx-prometheus-exporter 1.5.3', 'monitoring', 'Metrics Nginx :9113'],
    ['cadvisor', 'cadvisor 0.55.1', 'monitoring', 'Metrics container :8080'],
    ['prometheus', 'prom/prometheus v3.5.0', 'monitoring, management', 'TSDB, host 9090'],
    ['grafana', 'grafana/grafana 12.1.1', 'monitoring, management', 'Dashboard, host 3001'],
    ['loki', 'grafana/loki 3.5.1', 'monitoring', 'Lưu và truy vấn log :3100'],
    ['promtail', 'grafana/promtail 3.5.1', 'monitoring', 'Thu log từ 3 volume'],
], [3.0, 4.6, 3.6, 4.8], 'Danh sách 13 service trong compose.yaml')
para('Mọi cổng publish đều bind vào `127.0.0.1`, nên các giao diện chỉ truy cập được từ chính máy chạy Docker. '
     'Dữ liệu bền vững nằm trong chín named volume: `postgres_data`, `postgres_logs`, `pgadmin_data`, `app_logs`, '
     '`nginx_logs`, `prometheus_data`, `grafana_data`, `loki_data`, `promtail_positions`. Các image hạ tầng được '
     'ghim theo digest SHA-256 (`infra/images.lock.json`) để cố định phiên bản hạ tầng. Image app được build từ source và dependency lockfile; phiên bản đã chạy cần đối chiếu image ID.')

section('1.2. Yêu cầu quản lý mã nguồn trên GitHub')
para('Toàn bộ mã nguồn ứng dụng và file cấu hình (`compose.yaml`, `infra/nginx/default.conf`, '
     '`infra/prometheus/prometheus.yml`, `infra/loki/config.yaml`, `infra/promtail/config.yaml`, provisioning '
     'Grafana, Dockerfile, script vận hành) được quản lý trên repository GitHub đặt tên theo mã số sinh viên: '
     '**github.com/vzmanhdev/dtc245200223_dat_lich_hen**, nhánh `main`.')
table(['Thư mục / file', 'Nội dung'], [
    ['`app/`', 'Mã nguồn Express (`server.js`, `domain.js`, `management.js`, `customer-auth.js`), giao diện EJS, '
               'schema SQL v1–v3, `migrate.js`, Dockerfile'],
    ['`infra/nginx/`', '`default.conf` (reverse proxy, TLS, headers, rate limit), script xoay log'],
    ['`infra/prometheus/`', '`prometheus.yml` – 5 scrape job'],
    ['`infra/grafana/`', 'Datasource Prometheus/Loki và dashboard `appointment.json` (15 panel)'],
    ['`infra/loki/`, `infra/promtail/`', 'Cấu hình lưu trữ log và thu log'],
    ['`scripts/`', 'Tạo `.env`/chứng chỉ, khởi động, kiểm tra stack, TLS, backup/restore, rollback'],
    ['`tests/`', 'Unit test và integration test trên PostgreSQL thật'],
    ['`docs/`', 'Phân tích yêu cầu, runbook, hướng dẫn quản trị, LogQL, ảnh minh chứng'],
], [4.6, 11.4], 'Cấu trúc repository')
para('Lịch sử main đã được gộp thành đúng ba mốc Nginx, monitoring và logging/hardening. Đây là lịch sử tổ chức lại có sao lưu cục bộ, không phải khẳng định toàn bộ phần mềm được viết trong ba lần thao tác. Các bản sửa nghiệp vụ và vận hành được tổng hợp vào mốc cuối.')
table(['Mốc', 'Định danh', 'Nội dung và cách giải thích'],[
 ['Commit 1','`7e2bec8`','Website + PostgreSQL + pgAdmin + Compose và Nginx reverse proxy: đưa website lên và lưu dữ liệu.'],
 ['Commit 2','`5a8ee51`','Prometheus + Grafana + exporters/cAdvisor: quan sát container, web server và DB.'],
 ['Commit 3','`main / HEAD`','Loki + Promtail + ba LogQL; hoàn thiện hardening, HTTPS, kiểm thử và báo cáo. SHA cuối tra bằng git log -1.']
], [2.3,3.0,10.7], 'Ba mốc Git trên nhánh main')
figure(E / 'github-three-milestones.png','GitHub hiển thị đúng ba mốc trên main',15.0,9.0)
para('Ảnh tập trung vào tên mốc và ngữ cảnh repository; SHA của mốc cuối được tra trực tiếp trên main vì báo cáo nằm trong chính commit đó. Tên repo chứa MSSV, nhưng tài khoản `vzmanhdev` chưa đặt theo MSSV như nguyên văn đề; đây là điểm hình thức cần đối chiếu giảng viên.')
para('File `README.md` hướng dẫn đầy đủ: yêu cầu môi trường, clone repository, tạo file biến môi trường, khởi '
     'chạy toàn bộ stack, địa chỉ từng giao diện, kịch bản dùng thử nghiệp vụ và lệnh kiểm thử. Các bước chạy lần đầu:')
code('git clone https://github.com/vzmanhdev/dtc245200223_dat_lich_hen.git',
     'cd dtc245200223_dat_lich_hen',
     'npm ci --ignore-scripts',
     'npm run setup            # sinh .env với mật khẩu ngẫu nhiên + chứng chỉ TLS localhost',
     'docker compose up -d --build --wait --wait-timeout 240',
     'docker compose ps')
para('`npm run setup` tạo file `.env` chứa mật khẩu ngẫu nhiên riêng cho PostgreSQL, role ứng dụng, role giám sát, '
     'tài khoản admin, pgAdmin, Grafana và khóa session. File `.env`, khóa TLS và file backup đều nằm trong '
     '`.gitignore` nên không bao giờ được đưa lên GitHub.')

# ---------------------------------------------------------------- Chapter 2
current_chapter[0] = 2
chapter('CHƯƠNG 2: TRIỂN KHAI ỨNG DỤNG WEB, CƠ SỞ DỮ LIỆU VÀ REVERSE PROXY')
section('2.1. Triển khai ứng dụng Web & Cơ sở dữ liệu (Bước 1 & Bước 2)')
sub('a) Chức năng của website đặt lịch hẹn')
para('Mô hình nghiệp vụ: khách hàng đặt lịch tư vấn với một chuyên viên của đơn vị. Website đáp ứng đủ bốn nhóm '
     'chức năng mà đề 40 nêu:')
bullet('Dịch vụ', 'danh sách dịch vụ có thời lượng và giá; admin tạo, sửa, ngừng dịch vụ và phân công chuyên viên '
       'cho từng dịch vụ.')
bullet('Khung giờ', 'admin mở slot đơn lẻ hoặc hàng loạt cho từng chuyên viên, đóng/mở slot. Một chuyên viên không '
       'có hai slot chồng lấn; hai chuyên viên khác nhau có thể nhận khách cùng giờ.')
bullet('Khách hàng', 'khách đăng ký tài khoản (mật khẩu tối thiểu 12 ký tự), đăng nhập, cập nhật hồ sơ, xem "Lịch '
       'hẹn của tôi". Admin tra cứu lịch sử khách và khóa/mở tài khoản.')
bullet('Xác nhận', 'lịch mới ở trạng thái `pending`; chỉ admin được xác nhận (`confirmed`) hoặc hoàn tất '
       '(`completed`); khách được hủy trước giờ bắt đầu. Mọi thay đổi trạng thái được ghi vào `appointment_events`.')
figure(E / 'customer-booking-success.png', 'Khách hàng đặt lịch thành công qua website', 14.0, 11.0)
figure(E / 'admin-confirmation-desktop.png', 'Trang quản trị: tiếp nhận và xác nhận lịch hẹn', 15.0, 11.0)
sub('b) Cơ sở dữ liệu PostgreSQL')
para('Schema được áp dụng theo ba phiên bản bởi job `migrate` và lưu vết trong bảng `schema_migrations`. Các bảng '
     'chính: `services`, `staff`, `staff_services`, `time_slots`, `customer_accounts`, `customers`, `appointments`, '
     '`appointment_events`, `admin_users` và `session` (lưu phiên đăng nhập qua `connect-pg-simple`).')
figure(D / '04-actual-erd.png', 'Sơ đồ quan hệ dữ liệu (ERD) thực tế', 15.0, 15.0)
para('Tính toàn vẹn được đảm bảo ngay ở tầng cơ sở dữ liệu chứ không chỉ ở ứng dụng: unique index trên slot đang '
     'được giữ khiến hai request đồng thời đặt cùng một slot chỉ có một request thành công; khóa ngoại, ràng buộc '
     'CHECK và transaction bảo vệ lịch sử khách hàng. Đối tượng được "ngừng hoạt động" thay vì xóa cứng.')
figure(D / '05-appointment-states.png', 'Vòng đời trạng thái lịch hẹn', 13.0, 9.0)
para('Ứng dụng kết nối bằng ba role tách biệt theo nguyên tắc quyền tối thiểu:')
table(['Role', 'Dùng bởi', 'Quyền'], [
    ['`postgres`', 'migrate, pgAdmin', 'Chủ sở hữu schema; chỉ dùng khi tạo/nâng cấp schema và quản trị'],
    ['`appointment_app`', 'app', 'Chỉ SELECT/INSERT/UPDATE trên đúng bảng/cột cần thiết; '
                                 'không superuser, createdb, createrole'],
    ['`appointment_monitor`', 'postgres-exporter', 'Chỉ đọc thống kê để xuất metrics'],
], [4.0, 3.5, 8.5], 'Phân tách role PostgreSQL')
sub('c) Đóng gói container và file Compose')
para('Image ứng dụng được dựng bằng Dockerfile hai giai đoạn: giai đoạn đầu dùng `node:24-alpine3.23` để chạy '
     '`npm ci --omit=dev`; giai đoạn runtime chỉ chứa Alpine 3.23, binary Node và dependency production, chạy bằng '
     'user `node` (UID 1000) và có HEALTHCHECK gọi `/health/ready`. Trích đoạn `compose.yaml` cho tầng ứng dụng và '
     'cơ sở dữ liệu:')
code('services:',
     '  postgres:',
     '    image: postgres:17-bookworm@sha256:ec4dfc65...',
     '    environment:',
     '      POSTGRES_PASSWORD: ${POSTGRES_PASSWORD:?Run npm run setup}',
     '      POSTGRES_INITDB_ARGS: --auth-host=scram-sha-256',
     '    volumes: ["postgres_data:/var/lib/postgresql/data", "postgres_logs:/var/log/postgresql"]',
     '    networks: [db]',
     '    healthcheck: {test: ["CMD-SHELL", "pg_isready -U $${POSTGRES_USER} -d $${POSTGRES_DB}"]}',
     '  migrate:',
     '    command: ["node", "app/migrate.js"]',
     '    depends_on: {postgres: {condition: service_healthy}}',
     '  app:',
     '    image: ${APP_IMAGE:-manh-appointment:local}',
     '    networks: [frontend, db, monitoring]',
     '    read_only: true',
     '    cap_drop: [ALL]',
     '    depends_on: {migrate: {condition: service_completed_successfully}}',
     '  pgadmin:',
     '    image: dpage/pgadmin4:9@sha256:33eb7d07...',
     '    ports: ["127.0.0.1:${PGADMIN_PORT:-5050}:8080"]',
     '    networks: [db, management]',
     'networks:',
     '  db: {internal: true}')
para('Thứ tự khởi động được điều khiển bằng `depends_on` có điều kiện: `log-init` → `postgres` (healthy) → '
     '`migrate` (hoàn tất) → `app` (healthy) → `nginx`. Nhờ vậy lệnh `docker compose up -d --wait` chỉ báo thành '
     'công khi cả chuỗi đã sẵn sàng.')
sub('d) Công cụ quản lý DB – pgAdmin')
para('pgAdmin 4 được khai báo sẵn server "Appointment PostgreSQL" qua `infra/postgres/servers.json` (host '
     '`postgres`, cổng 5432, database `appointment`). Người quản trị đăng nhập bằng tài khoản trong `.env`, nhập '
     'mật khẩu PostgreSQL khi kết nối, rồi duyệt bảng hoặc chạy truy vấn SQL trên dữ liệu thật. pgAdmin bật '
     '`ENHANCED_COOKIE_PROTECTION` và yêu cầu master password.')
figure(E / 'pgadmin-services-verified.png', 'pgAdmin truy vấn trực tiếp bảng dữ liệu của ứng dụng', 15.0, 11.0)

section('2.2. Cấu hình Nginx làm Reverse Proxy và Bảo mật cơ bản (Bước 3 - Commit 1)')
para('Nginx là điểm vào duy nhất của website. Container app không publish cổng nào ra host; mọi request phải đi '
     'qua Nginx trên mạng `frontend`. Nginx dùng Docker DNS resolver (`127.0.0.11`) để luôn tìm đúng IP của app '
     'kể cả sau khi container app được tạo lại.')
para('Nginx giới hạn 5 request/giây/IP (burst 40), body 16 KB; chặn public /metrics, /health/* và /stub_status. Exporter đọc cổng nội bộ 8081. Log JSON ghi method, status, thời lượng và request_id, không ghi body, token hoặc query string.')
figure(D / '01-web-data-flow.png', 'Luồng request web: trình duyệt → Nginx → app → PostgreSQL', 15.0, 10.0)
para('Trích đoạn cấu hình `infra/nginx/default.conf`:')
code('limit_req_zone $binary_remote_addr zone=booking:10m rate=5r/s;',
     'server {',
     '    listen 8080;',
     '    listen 8443 ssl;',
     '    ssl_certificate     /etc/nginx/certs/localhost.crt;',
     '    ssl_certificate_key /etc/nginx/certs/localhost.key;',
     '    ssl_protocols TLSv1.2 TLSv1.3;',
     '    server_tokens off;',
     '    client_max_body_size 16k;',
     '    location = /metrics { return 404; }',
     '    location ^~ /health/ { return 404; }',
     '    location / {',
     '        limit_req zone=booking burst=40 nodelay;',
     '        limit_req_status 429;',
     '        resolver 127.0.0.11 valid=5s ipv6=off;',
     '        set $appointment_backend app:3000;',
     '        proxy_pass http://$appointment_backend;',
     '        proxy_set_header X-Forwarded-Proto $scheme;',
     '        proxy_hide_header X-Powered-By;',
     '    }',
     '}')
para('**HTTPS tự ký:** `scripts/setup-tls.js` tạo chứng chỉ tự ký cho `localhost`/`127.0.0.1` (có SAN). '
     'Nginx chỉ chấp nhận TLS 1.2 và 1.3. Kiểm tra bằng `scripts/verify-tls.js` cho kết quả: HTTP 200, giao thức '
     '**TLSv1.3**, chứng chỉ được xác thực với CA cục bộ, cookie phiên mang cờ `Secure` và các endpoint nội bộ bị '
     'chặn. Do là chứng chỉ tự ký, trình duyệt cần tin cậy CA cục bộ hoặc chấp nhận cảnh báo; HSTS được tắt có chủ '
     'đích để không "khóa" trình duyệt vào chứng chỉ thử nghiệm.')
para('Host HTTPS_PORT=443 được Docker ánh xạ vào Nginx 8443, nên URL là https://localhost. HTTP 8080 là kênh lab phụ; Grafana/Prometheus/pgAdmin vẫn dùng HTTP loopback trên cổng riêng.')
figure(E / 'https-current-443.png', 'Truy cập website qua HTTPS https://localhost', 13.0, 9.0)
para('Các security headers thực tế trên phản hồi, gồm phần Nginx thêm vào và phần ứng dụng thêm bằng thư viện '
     '`helmet`:')
table(['Header', 'Giá trị', 'Tác dụng'], [
    ['X-Frame-Options', 'SAMEORIGIN', 'Chống Clickjacking (nhúng trang vào iframe lạ)'],
    ['X-Content-Type-Options', 'nosniff', 'Chống MIME-sniffing'],
    ['Content-Security-Policy', "default-src 'self'; ...", 'Chỉ cho phép script/style cùng nguồn – lớp chống XSS chính'],
    ['X-XSS-Protection', '0', 'Tắt bộ lọc XSS cũ của trình duyệt theo khuyến nghị hiện hành; '
                               'thay bằng CSP (giá trị "1; mode=block" đã lỗi thời)'],
    ['Referrer-Policy', 'strict-origin-when-cross-origin', 'Không lộ đường dẫn đầy đủ sang site khác'],
    ['Permissions-Policy', 'camera=(), microphone=(), geolocation=()', 'Vô hiệu API thiết bị không dùng đến'],
    ['Server', 'nginx (ẩn phiên bản)', '`server_tokens off`, ẩn `X-Powered-By` của Express'],
], [4.2, 4.6, 7.2], 'Security headers trên website')


# ---------------------------------------------------------------- Chapter 3
current_chapter[0] = 3
chapter('CHƯƠNG 3: GIÁM SÁT HỆ THỐNG VỚI PROMETHEUS, GRAFANA VÀ LOKI')
section('3.1. Tích hợp Prometheus + Grafana (Bước 4 - Commit 2)')
para('Prometheus hoạt động theo cơ chế **pull**: định kỳ 15 giây gọi HTTP tới endpoint metrics của từng target '
     'trên mạng `monitoring` nội bộ. Dữ liệu được lưu trong TSDB với giới hạn 7 ngày hoặc 1 GB.')
figure(D / '02-metrics-pull-flow.png', 'Luồng thu thập metrics: Prometheus scrape, Grafana truy vấn PromQL', 15.0, 10.0)
code('global:',
     '  scrape_interval: 15s',
     'scrape_configs:',
     '  - job_name: prometheus   # targets: ["prometheus:9090"]',
     '  - job_name: appointment  # targets: ["app:3000"]        – prom-client',
     '  - job_name: nginx        # targets: ["nginx-exporter:9113"]',
     '  - job_name: postgres     # targets: ["postgres-exporter:9187"]',
     '  - job_name: cadvisor     # targets: ["cadvisor:8080"]')
table(['Job', 'Nguồn metrics', 'Chỉ số tiêu biểu'], [
    ['appointment', 'Endpoint `/metrics` của app (thư viện prom-client)', 'Số request theo mã HTTP, độ trễ p95, '
                                                                         'CPU/heap của tiến trình Node'],
    ['nginx', 'nginx-exporter đọc `nginx:8081/stub_status`', '`nginx_up`, kết nối đang mở, request/giây'],
    ['postgres', 'postgres-exporter với role `appointment_monitor`', '`pg_up`, số kết nối, transaction commit/giây'],
    ['cadvisor', 'cgroup của Docker trên host', 'CPU, RAM working set, Network I/O theo từng container'],
    ['prometheus', 'Chính Prometheus', 'Tình trạng scrape, TSDB'],
], [2.6, 6.0, 7.4], 'Năm scrape job của Prometheus')
para('Grafana được cấu hình hoàn toàn bằng file (provisioning): datasource Prometheus và Loki trong '
     '`infra/grafana/provisioning/datasources`, dashboard **Appointment Localhost Overview** trong '
     '`infra/grafana/dashboards/appointment.json`. Dashboard có **15 panel**, tự làm mới mỗi 15 giây:')
table(['Nhóm', 'Panel'], [
    ['Tình trạng', 'Scrape targets UP; PostgreSQL connected; Nginx connected'],
    ['Web server', 'HTTP requests/second; Application requests by HTTP status; Application p95 response time; '
                   'Nginx active connections; Nginx requests/second'],
    ['Container', 'Container CPU (by Compose service); Container memory working set; Container network receive; '
                  'Container network transmit'],
    ['Database', 'PostgreSQL connections; PostgreSQL committed transactions/second'],
    ['Log', 'Appointment application logs (truy vấn Loki)'],
], [3.0, 13.0], 'Các panel của dashboard Grafana')
figure(E / 'grafana-dashboard-verified.png', 'Dashboard Grafana giám sát container, web server và database', 15.5, 11.0)
figure(E / 'grafana-network-verified.png', 'Panel Network I/O theo từng container', 15.5, 10.0)
para('Kết quả kiểm tra: cả 5 job ở trạng thái UP, `pg_up = 1` và `nginx_up = 1` (exporter đọc được dịch vụ nguồn), '
     'metrics CPU/RAM được gắn nhãn đúng service Compose. Grafana tắt đăng ký tài khoản, tắt truy cập ẩn danh và '
     'chỉ mở trên `127.0.0.1:3001`.')

section('3.2. Triển khai hệ thống Log tập trung Loki + Promtail (Bước 5 - Commit 3)')
para('Ứng dụng (thư viện `pino`) và Nginx ghi log dạng JSON, PostgreSQL ghi log text có xoay vòng theo ngày/10 MB. '
     'Ba nguồn này nằm trong ba volume `app_logs`, `nginx_logs`, `postgres_logs`. Promtail mount các volume ở chế '
     'độ **chỉ đọc**, lưu vị trí đã đọc vào `promtail_positions` để không gửi trùng khi khởi động lại, và đẩy log '
     'qua HTTP tới `loki:3100`. Loki lưu chỉ mục TSDB v13 trên filesystem với thời gian giữ 168 giờ.')
figure(D / '03-log-collection-flow.png', 'Thu log bằng Promtail Loki và tra cứu từ Grafana', 15.0, 10.0)
code('scrape_configs:',
     '  - job_name: appointment',
     '    pipeline_stages:',
     '      - json: {expressions: {event_time: time}}',
     '      - timestamp: {source: event_time, format: RFC3339Nano}',
     '    static_configs:',
     '      - labels: {service: app, environment: local,',
     '                 __path__: /var/log/appointment/app.log}',
     '  - job_name: nginx      # service: nginx,    /var/log/appointment-nginx/access.log',
     '  - job_name: postgres   # service: postgres, /var/log/appointment-postgres/*.log')
para('Các truy vấn LogQL đã chạy và có kết quả trong Grafana Explore (datasource Loki):')
bullet('Truy vấn 1 (Toàn bộ log của ứng dụng)', '')
code('{service="app", environment="local"}')
bullet('Truy vấn 2 (Lọc request lỗi HTTP ≥ 400 từ Nginx access log, parse JSON)', '')
code('{service="nginx", environment="local"} | json | __error__="" | status >= 400')
bullet('Truy vấn 3 (Đếm số request lỗi trong cửa sổ 5 phút – dạng metric)', '')
code('sum(count_over_time({service="nginx", environment="local"} | json | __error__="" | status >= 400 [5m]))')
bullet('Truy vấn bổ sung (Log lỗi của ứng dụng, Pino level 50 = error; log PostgreSQL)', '')
code('{service="app", environment="local"} | json | __error__="" | level=50',
     '{service="postgres", environment="local"}')
figure(E / 'loki-logql-verified.png', 'Grafana Explore chạy LogQL lọc request lỗi từ Nginx', 15.5, 10.0)
para('Bộ kiểm tra `scripts/verify-stack.js` gửi cả ba truy vấn chính tới Loki và xác nhận mỗi truy vấn trả về ít '
     'nhất một stream. Log container được giới hạn bởi driver `json-file` (10 MB × 3 file); log Nginx được xoay '
     'bằng script đi kèm container, Promtail chỉ đọc file hiện hành để tránh thu lại file xoay; timestamp sự kiện JSON và positions được xử lý riêng, log trùng lịch sử trước bản sửa vẫn theo retention.')

# ---------------------------------------------------------------- Chapter 4
current_chapter[0] = 4
chapter('CHƯƠNG 4: CÁC BIỆN PHÁP HARDENING VÀ ĐÁNH GIÁ BẢO MẬT')
section('4.1. Áp dụng các biện pháp Hardening (Bước 6)')
para('Các biện pháp gia cố được áp dụng ở cả tầng container, mạng, cơ sở dữ liệu và ứng dụng:')
bullet('Chạy container dưới quyền non-root user', 'image app tạo user `node` (UID 1000) và khai báo `USER node`; '
       'Nginx dùng image `nginx-unprivileged` (UID 101) lắng nghe cổng 8080/8443 không cần root. Kết hợp '
       '`no-new-privileges:true` cho phần lớn service để hạn chế việc nhận thêm đặc quyền khi thực thi.')
bullet('Hệ thống file chỉ đọc và bỏ capability', 'app, migrate, nginx, các exporter và promtail chạy '
       '`read_only: true`; app và Nginx có tmpfs `/tmp` với `noexec,nosuid`, dung lượng 16 MB. `cap_drop: [ALL]` loại bỏ toàn bộ Linux '
       'capability trên hầu hết service.')
bullet('Phân tách mạng nội bộ (Network Isolation)', 'bốn network: `frontend` (nginx ↔ app), `db` internal '
       '(chỉ app, migrate, pgAdmin, postgres-exporter thấy PostgreSQL), `monitoring` internal và `management` cho '
       'các UI quản trị. PostgreSQL, Loki, exporter không có cổng trên host; các cổng còn lại chỉ bind `127.0.0.1`.')
bullet('Mật khẩu mạnh và biến môi trường', '`npm run setup` sinh mật khẩu ngẫu nhiên riêng cho từng tài khoản '
       'vào `.env` (bị `.gitignore`); Compose dùng cú pháp `${VAR:?}` để từ chối khởi động nếu thiếu bí mật. '
       'PostgreSQL dùng `scram-sha-256`; mật khẩu người dùng ứng dụng được băm trước khi lưu.')
bullet('Hạn chế quyền cơ sở dữ liệu', 'role `appointment_app` chỉ có quyền trên đúng bảng/cột cần dùng, không '
       'superuser/createdb/createrole; exporter dùng role giám sát riêng.')
bullet('Bảo mật ứng dụng', 'session lưu trong PostgreSQL, cookie `HttpOnly`/`SameSite`/`Secure` qua HTTPS, chống '
       'CSRF, rate limit đăng nhập, đổi session ID sau xác thực, vai trò do server quyết định. Khách không xem/hủy '
       'được lịch của người khác bằng cách đổi ID.')
bullet('Security headers và giới hạn tài nguyên', 'headers như mục 2.2; `mem_limit` cho từng service; image ghim '
       'digest; log có giới hạn dung lượng.')
bullet('Tường lửa', 'máy host là Windows nên dùng Windows Defender Firewall thay cho iptables/UFW; cả ba profile '
       'Domain/Private/Public đang bật. Các cổng UI bind loopback theo cấu hình; chưa nghiệm thu truy cập từ máy LAN khác.')
table(['Service', 'User', 'read_only', 'cap_drop ALL', 'Ghi chú'], [
    ['app / migrate', '1000 (node)', 'Có', 'Có', 'App có tmpfs /tmp; migrate là job'],
    ['nginx', '101 (nginx)', 'Có', 'Có', 'image unprivileged'],
    ['exporters', 'non-root (image)', 'Có', 'Có', ''],
    ['prometheus, loki, grafana', 'non-root (image)', 'Không', 'Có', 'cần ghi volume dữ liệu'],
    ['postgres, pgadmin', 'user của image', 'Không', 'Không', 'no-new-privileges, mạng db internal'],
    ['promtail', 'root', 'Có', 'Có, chỉ thêm DAC_READ_SEARCH', 'Ngoại lệ: đọc log của UID khác, mount chỉ đọc'],
    ['cadvisor', 'root', 'Không', 'Không', 'Ngoại lệ: privileged để đọc cgroup host'],
    ['log-init', 'root', '–', '–', 'Job một lần, không có mạng'],
], [3.6, 2.6, 2.4, 3.2, 4.2], 'Tổng hợp hardening theo từng service')
para('Hai ngoại lệ `promtail` và `cadvisor` được giữ có chủ đích vì bản chất công cụ cần đọc dữ liệu của host; '
     'chúng nằm trong mạng `monitoring` internal, không publish cổng và không chia sẻ quyền này với ứng dụng.')

section('4.2. Đánh giá bảo mật và kiểm thử')
para('Image ứng dụng được quét bằng **Docker Scout**. Lần quét đầu phát hiện 13 lỗ hổng mức High nằm trong các gói '
     'Debian và npm đi kèm base image cũ. Sau khi chuyển runtime sang Alpine 3.23 tối giản (không chứa npm/Yarn), '
     'quét lại không lọc mức độ:')
table(['Mức độ', 'Trước', 'Sau'], [
    ['Critical', '0', '0'], ['High', '13', '**0**'], ['Medium', 'chưa đo', '1'],
    ['Low / Unspecified', 'chưa đo', '0 / 0'], ['Số package nhận diện', '390', '141'],
], [6.0, 4.0, 4.0], 'Kết quả quét Docker Scout image ứng dụng')
para('Cảnh báo còn lại là **CVE-2025-60876** (BusyBox 1.37.0-r30, mức Medium) chưa có bản vá trên nhánh Alpine '
     '3.23 tại thời điểm quét; được ghi nhận để theo dõi, không dùng bộ lọc để che giấu.')
table(['Bộ kiểm tra', 'Kết quả'], [
    ['Unit test (`npm test`)', '4/4 đạt'],
    ['Integration test trên PostgreSQL thật', '3/3 nhóm đạt – đăng ký/đăng nhập, quyền sở hữu lịch, hai request '
                                             'tranh cùng slot, CSRF, quyền admin'],
    ['`scripts/verify-stack.js`', '16/16 đạt – engine, config, non-root, DB không publish, role ít quyền, headers, '
                                  'pgAdmin, 5 target, pg_up/nginx_up, dashboard 15 panel, 3 LogQL'],
    ['`scripts/verify-tls.js`', 'TLSv1.3, chứng chỉ hợp lệ với CA cục bộ, cookie Secure, endpoint nội bộ bị chặn'],
    ['Backup / restore', '`backup.ps1` tạo dump, `restore-test.ps1` khôi phục sang DB riêng và đối chiếu số dòng'],
    ['Clone sạch', 'Clone repository vào thư mục mới, dựng đủ 13 service và đạt các kiểm tra trên'],
], [5.5, 10.5], 'Kết quả kiểm thử và nghiệm thu')

sub('Kịch bản demo và đối chiếu yêu cầu')
para('Demo khoảng 15 phút: (1) mở https://localhost, đăng nhập khách, chọn dịch vụ/chuyên viên/slot và gửi yêu cầu; (2) admin xác nhận cùng mã lịch, khách xem lại trạng thái; (3) pgAdmin truy vấn dữ liệu tương ứng; (4) Prometheus Targets đối chiếu năm job, Grafana chỉ ra web/DB/CPU/RAM/Network I/O; (5) Explore chọn Loki và thực thi ba LogQL; (6) kiểm tra UID, quyền DB, network/headers và ba mốc Git. Không hiển thị .env hoặc mật khẩu lên màn chiếu.')
para('Bảy tiêu chí được đối chiếu như sau: GitHub ở mục 1.2; ứng dụng/DB ở 2.1; Nginx ở 2.2; monitoring ở 3.1; log ở 3.2; hardening ở 4.1; kiểm thử/demo ở 4.2. Báo cáo mô tả kết quả kiểm chứng, không tự chấm điểm thay giảng viên.')
# ---------------------------------------------------------------- Conclusion
current_chapter[0] = 0
chapter('KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN')
para('Báo cáo đã trình bày quy trình triển khai hệ thống đặt lịch hẹn theo đúng 6 bước của đề 40. Kết quả đạt '
     'được:')
table(['Bước', 'Yêu cầu', 'Kết quả'], [
    ['1', 'Quản lý mã nguồn trên GitHub', 'Repository dtc245200223_dat_lich_hen, README chi tiết, commit theo tiến độ'],
    ['2', 'Ứng dụng + DB + công cụ quản lý', 'Express/EJS + PostgreSQL 17 + pgAdmin 4, đủ dịch vụ/khung giờ/khách '
                                            'hàng/xác nhận'],
    ['3', 'Nginx reverse proxy (Commit 1)', 'HTTPS tự ký TLS 1.2/1.3, security headers, rate limit'],
    ['4', 'Prometheus + Grafana (Commit 2)', '5 target UP, dashboard 15 panel cho container/web/DB'],
    ['5', 'Loki + Promtail (Commit 3)', '3 nguồn log và 3 query chính đã kiểm chứng; query lỗi app chỉ có dữ liệu khi phát sinh lỗi'],
    ['6', 'Hardening', 'Non-root, read-only, cap_drop, mạng internal, mật khẩu ngẫu nhiên, role DB tối thiểu, '
                      'image không còn lỗ hổng High'],
], [1.4, 5.0, 9.6], 'Tổng hợp kết quả 6 bước')
para('Hệ thống khởi chạy hoàn chỉnh bằng một lệnh `docker compose up`, có thể dựng lại từ clone sạch, có giám sát '
     'metrics, log tập trung và các biện pháp bảo mật được kiểm chứng bằng script tự động.')
para('**Hạn chế:** hệ thống mới triển khai trên localhost, chưa đưa lên máy chủ Internet; chứng chỉ TLS tự ký; '
     'Promtail chỉ thu ba nguồn log cốt lõi; chưa cấu hình Alertmanager; chưa có tài khoản đăng nhập riêng cho '
     'chuyên viên, gửi email/SMS nhắc lịch, quên mật khẩu hay thanh toán.')
para('**Hướng phát triển:** triển khai lên VPS với tên miền và chứng chỉ Let\'s Encrypt; bổ sung Alertmanager gửi '
     'cảnh báo; tự động hóa build, quét image và deploy bằng GitHub Actions; tích hợp hệ thống phát hiện xâm nhập '
     'như Falco; chuyển từ Promtail sang Grafana Alloy; mở rộng nghiệp vụ tài khoản chuyên viên và thông báo email.')

chapter('TÀI LIỆU THAM KHẢO')
for ref in [
 '[1] Đề tài Thực hành môn Triển khai và Quản trị Hệ thống Phần mềm, yêu cầu chung và Đề 40. Tài liệu học phần do sinh viên cung cấp.',
 '[2] Bao_Cao_Trien_Khai_He_Thong.docx. Khung bố cục bốn chương được giữ; nội dung và minh chứng điều chỉnh theo cấu hình thật.',
 '[3] Docker Docs. Networking in Compose. https://docs.docker.com/compose/how-tos/networking/',
 '[4] Nginx Documentation. Configuring HTTPS servers. https://nginx.org/en/docs/http/configuring_https_servers.html',
 '[5] PostgreSQL 17 Documentation. https://www.postgresql.org/docs/17/',
 '[6] Prometheus Documentation. https://prometheus.io/docs/introduction/overview/',
 '[7] Grafana Loki Documentation. Promtail agent. https://grafana.com/docs/loki/latest/send-data/promtail/',
 '[8] Repository source/config và minh chứng thực tế. https://github.com/vzmanhdev/dtc245200223_dat_lich_hen. Đối chiếu tài liệu ngày 25/10/2026.'
]:para(ref,11.5,after=100)
settings=d.settings.element
u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
d.core_properties.author='Nguyễn Văn Mạnh'
d.core_properties.last_modified_by='Nguyễn Văn Mạnh'
d.core_properties.title='Báo cáo triển khai hệ thống đặt lịch hẹn theo khung bốn chương'
d.core_properties.subject='Đề 40 dtc245200223 CNTTK23B'
qa=ROOT/'tmp/report-compact';qa.mkdir(parents=True,exist_ok=True)
(qa/'manifest.json').write_text(json.dumps({'output':str(OUT),'reference':str(TEMPLATE),'source_sha256':sha256(TEMPLATE.read_bytes()).hexdigest(),'figures':figure_titles,'tables':table_titles,'main_commits':3},ensure_ascii=False,indent=2),encoding='utf8')
OUT.parent.mkdir(parents=True, exist_ok=True)
# Restore the approved institutional cover independently of the body template.
cover_reference = ROOT / 'docx/Bao_cao_Appointment_System.docx'
cover_doc = Document(cover_reference)
set_text(cover_doc.paragraphs[7], '      SINH VIÊN THỰC HIỆN:')
set_text(cover_doc.paragraphs[8], '      HỌ VÀ TÊN: NGUYỄN VĂN MẠNH     LỚP: CNTTK23B\n      MÃ SINH VIÊN: dtc245200223\n\nGIẢNG VIÊN HƯỚNG DẪN: NGUYỄN ANH CHUYÊN')
for run in cover_doc.paragraphs[8].runs:
    run.font.size = Pt(13)
set_text(cover_doc.paragraphs[14], 'Thái Nguyên, ngày 25 tháng 10 năm 2026')
cover_elements = [deepcopy(el) for el in list(cover_doc.element.body)[:15]]
for el in cover_elements:
    for node in el.iter():
        for attr in (qn('r:embed'), qn('r:id'), qn('r:link')):
            rid = node.get(attr)
            if rid and rid in cover_doc.part.rels:
                rel = cover_doc.part.rels[rid]
                if rel.reltype.endswith('/image'):
                    new_rid, _ = d.part.get_or_add_image(BytesIO(rel.target_part.blob))
                    node.set(attr, new_rid)
    for sp in el.iter(qn('w:sectPr')):
        for child in list(sp):
            if child.tag in (qn('w:headerReference'), qn('w:footerReference')):
                sp.remove(child)
        title = OxmlElement('w:titlePg'); sp.append(title)
# Existing first nine paragraphs are the temporary generic cover.
for el in list(d.element.body)[:9]:
    d.element.body.remove(el)
for index, el in enumerate(cover_elements):
    d.element.body.insert(index, el)
# Preserve Letter geometry and the existing footer for all body pages.
body_section = d.element.body.find(qn('w:sectPr'))
for title in list(body_section.findall(qn('w:titlePg'))):
    body_section.remove(title)

d.save(OUT)
print(OUT)
