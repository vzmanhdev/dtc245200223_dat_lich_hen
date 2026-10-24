"""Detailed academic report with live Word TOC, figure/table lists and traceable evidence.
Requires python-docx/Pillow from the document runtime; original sources stay unchanged.
"""
from pathlib import Path
from hashlib import sha256
from datetime import datetime, timezone, timedelta
import subprocess, json
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent.parent
REF=ROOT/'docx/Bao_cao_Appointment_System.docx'
OUT=ROOT/'reports/Bao_cao_Trien_khai_Appointment_dtc245200223_Toan_bo_he_thong.docx'
QA=ROOT/'tmp/report-v2';QA.mkdir(parents=True,exist_ok=True)
E=ROOT/'docs/evidence';D=E/'diagrams'
GIT='C:/Users/ADMIN/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/git/cmd/git.exe'
raw=subprocess.check_output([GIT,'log','--reverse','--format=%H%x1f%h%x1f%aI%x1f%an%x1f%s'],cwd=ROOT,encoding='utf8')
commits=[r.split('\x1f') for r in raw.splitlines()]
head=commits[-1][0]
(QA/'artifact.md').write_text(f'''# Detailed report contract
Reference: {REF}; SHA256 {sha256(REF.read_bytes()).hexdigest()}.
Reference already rendered by Word: tmp/report-reference/reference.pdf, 35 pages.
User specifically retains cover layout (page 1), requests a complete, more detailed
academic report using both supplied sample structures, separate contents and figures.
Cover paragraphs 0..14 retained including logo, double page border and section break;
student name/MSSV/class/date replaced; instructor retained from source with provenance.
Body of other student's report removed entirely. Both source DOCX files remain read-only.
A4, cover original margins; body left3/right2/top2.35/bottom2.2cm, Times New Roman13,
1.3 line spacing; Heading1 16pt/Heading2 14pt/Heading3 13pt, black. Body justified.
Automatic TOC levels1..2; separate TOC\\c Hinh and TOC\\c Bang fields; sequential
captions, original PAGE footer. Word must update all fields/TOCs, repaginate and export.
Six accurate diagrams from source; full-system overview, separate web/metrics/log flows, actual ERD,
appointment states. Source-derived screenshots carry sources and explanatory notes.
Git table is all {len(commits)} commits at {head}; snapshot excludes later report commit.
Retain cover/footer/style relationship system; remove unused old sample image links.
Write a new filename because previous report is open (Office lock exists).
Render exact final DOCX with Word backend + render_docx.py/Poppler; inspect all pages.
''',encoding='utf8')

d=Document(REF)
cover=list(d.element.body)[:15]
for el in list(d.element.body):
    if el not in cover and el.tag!=qn('w:sectPr'):d.element.body.remove(el)
def cover_text(index,text):
    p=d.paragraphs[index];p.runs[0].text=text
    for r in p.runs[1:]:r.text=''
cover_text(7,'      SINH VIÊN THỰC HIỆN:')
cover_text(8,'      HỌ VÀ TÊN: NGUYỄN VĂN MẠNH     LỚP: CNTTK23B\n      MÃ SINH VIÊN: dtc245200223\n\nGIẢNG VIÊN HƯỚNG DẪN: NGUYỄN ANH CHUYÊN')
for r in d.paragraphs[8].runs:r.font.size=Pt(13)
cover_text(14,'Thái Nguyên, ngày 24 tháng 10 năm 2026')
d.sections[1].different_first_page_header_footer=False
def style(name,size=13,bold=False):
    st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    rfonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
    if rfonts is not None:
        for attr in list(rfonts.attrib):
            if 'Theme' in attr:del rfonts.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:rfonts.set(qn('w:'+attr),'Times New Roman')
    pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.3;pf.left_indent=Cm(0);pf.first_line_indent=Cm(0)
    pf.page_break_before=False;pf.widow_control=True;return st
body=style('Report Body');body.paragraph_format.first_line_indent=Cm(.6);body.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for n,size in [('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
    st=style(n,size,True);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.line_spacing=1.15
front=style('Report Front',16,True);front.paragraph_format.keep_with_next=True;front.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['Report Caption','Report Table Caption']:
    st=style(name,11.5,True);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;st.paragraph_format.line_spacing=1.1
    st.paragraph_format.keep_with_next=(name=='Report Table Caption')
style('Report Note',11);d.styles['Report Note'].paragraph_format.line_spacing=1.1
style('Report Code',9.5);d.styles['Report Code'].paragraph_format.line_spacing=1.05
for name in ['TOC 1','TOC 2','TOC 3','Table of Figures']:
    st=style(name,12);st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_after=Pt(5)
    if name=='TOC 1':st.font.bold=True
d.core_properties.author='Nguyễn Văn Mạnh';d.core_properties.last_modified_by='Nguyễn Văn Mạnh'
d.core_properties.title='Báo cáo triển khai và quản trị hệ thống đặt lịch hẹn'
d.core_properties.subject='Đề 40 – Nguyễn Văn Mạnh – dtc245200223 – CNTTK23B'
d.core_properties.comments='Bản chi tiết có mục lục, danh mục hình/bảng và minh chứng thực tế.'
figures=[];tables=[]
def field(p,instruction,placeholder=''):
    r=p.add_run();start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');r._r.append(start)
    r=p.add_run();ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' '+instruction+' ';r._r.append(ins)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    if placeholder:p.add_run(str(placeholder))
    r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
def front_page(title,first=False):
    q=d.add_paragraph(title,'Report Front');q.paragraph_format.page_break_before=not first
def chapter(title):
    q=d.add_paragraph(title,'Heading 1');q.paragraph_format.page_break_before=True
def h(title):d.add_paragraph(title,'Heading 2')
def h3(title):d.add_paragraph(title,'Heading 3')
def p(text):return d.add_paragraph(text,'Report Body')
def note(text):return d.add_paragraph(text,'Report Note')
def code(text):
    q=d.add_paragraph(style='Report Code');r=q.add_run(text);r.font.name='Consolas';r.font.size=Pt(9.5)
    q.paragraph_format.keep_together=True
    return q
def caption(label,title):
    seq='Hinh' if label=='Hình' else 'Bang';items=figures if label=='Hình' else tables;items.append(title)
    q=d.add_paragraph(style='Report Caption' if label=='Hình' else 'Report Table Caption')
    q.add_run(label+' ');field(q,f'SEQ {seq} \\* ARABIC',len(items));q.add_run('. '+title)
def table(title,headers,rows,widths=None):
    caption('Bảng',title);t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [16/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'888888');borders.append(b)
    t._tbl.tblPr.append(borders)
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');c._tc.get_or_add_tcPr().append(sh)
            for q in c.paragraphs:
                q.paragraph_format.space_before=Pt(3);q.paragraph_format.space_after=Pt(4);q.paragraph_format.line_spacing=1.1
                for r in q.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.font.bold=(i==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def fig(path,title,explain,source=None,width=16,max_height=14.5):
    path=Path(path);w,h=Image.open(path).size
    width=min(width,max_height*w/h)
    q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True;q.paragraph_format.space_after=Pt(3)
    q.add_run().add_picture(str(path),width=Cm(width))
    caption('Hình',title)
    note('Nguồn: '+(source or ('tác giả xây dựng theo source/config ngày 24/10/2026' if path.parent==D else 'ảnh chụp hệ thống thực hành ngày 24/10/2026'))+'.')
    p(explain)

front_page('LỜI MỞ ĐẦU',True)
p('Trong học phần Triển khai và Quản trị Hệ thống Phần mềm, việc hoàn thành một website mới chỉ là bước đầu. Một hệ thống có thể sử dụng được còn cần môi trường chạy tái lập, dữ liệu bền vững, cổng truy cập được bảo vệ, khả năng quan sát và quy trình xử lý sự cố. Vì vậy, em lựa chọn đề tài số 40 – Hệ thống đặt lịch hẹn – để thực hành đồng thời các nội dung xây dựng ứng dụng và tổ chức hạ tầng vận hành.')
p('Bài toán được cụ thể hóa thành một đơn vị tư vấn có nhiều chuyên viên. Khách hàng có tài khoản riêng, chọn dịch vụ, chuyên viên và khung giờ; quản trị viên tiếp nhận, xác nhận và theo dõi lịch. Hệ thống dùng frontend EJS/JavaScript, backend Express và PostgreSQL. Toàn bộ thành phần được triển khai bằng Docker Compose trên máy cá nhân Windows/WSL 2, với Nginx hỗ trợ HTTPS tự ký, Prometheus/Grafana thu thập metrics và Loki/Promtail tập trung log.')
p('Báo cáo được tổ chức theo trình tự: xác định yêu cầu và kiến trúc; triển khai web, cơ sở dữ liệu, reverse proxy; xây dựng monitoring và logging; áp dụng hardening; cuối cùng kiểm thử và đánh giá. Từng thành phần được trình bày theo mục đích, cách cấu hình, tác dụng và bằng chứng kiểm tra, thay vì chỉ liệt kê tên công cụ. Lịch sử Git, sơ đồ dữ liệu, cấu hình quan trọng và kịch bản demo được đưa vào để người đọc có thể đối chiếu với source thực tế.')
p('Các ảnh kết quả thuộc môi trường thực hành của Nguyễn Văn Mạnh; số dòng và giá trị dashboard phản ánh thời điểm chụp. Báo cáo không nhận kết quả của tài liệu mẫu là của hệ thống hiện tại. Phạm vi được thống nhất là localhost, chưa triển khai VPS hay Internet. Các giới hạn như chứng chỉ tự ký, tài khoản nhân viên chưa tách riêng và cảnh báo bảo mật còn lại được trình bày cụ thể ở phần đánh giá.')
p('Em xin cảm ơn giảng viên và học phần đã cung cấp kiến thức, đề bài và tài liệu thực hành để em vận dụng vào quá trình xây dựng, triển khai và kiểm tra hệ thống. Báo cáo là cơ sở để trình bày kết quả và tiếp tục hoàn thiện kỹ năng vận hành phần mềm.')
note('Sinh viên thực hiện: Nguyễn Văn Mạnh – MSSV dtc245200223 – lớp CNTTK23B.')
front_page('MỤC LỤC');field(d.add_paragraph(),r'TOC \o "1-2" \h \z \u')
front_page('DANH MỤC HÌNH ẢNH');field(d.add_paragraph(),r'TOC \h \z \c "Hinh"')
front_page('DANH MỤC BẢNG');field(d.add_paragraph(),r'TOC \h \z \c "Bang"')
front_page('DANH MỤC TỪ VIẾT TẮT')
table('Các thuật ngữ sử dụng trong báo cáo',['Từ viết tắt','Nghĩa và cách hiểu trong đề tài'],[
('API','Giao diện để frontend hoặc công cụ kiểm thử gọi chức năng backend'),('DB / SQL','Cơ sở dữ liệu / ngôn ngữ truy vấn có cấu trúc'),('HTTP / HTTPS','Giao thức web / HTTP được bảo vệ bằng TLS'),('TLS / SAN','Mã hóa kết nối / tên DNS và IP hợp lệ trong chứng chỉ'),('UID / PID','Định danh người dùng / tiến trình trong môi trường chạy'),('CSRF / XSS','Giả mạo yêu cầu từ trình duyệt / chèn mã script vào nội dung'),('PK / FK','Khóa chính / khóa ngoại trong mô hình dữ liệu'),('CRUD','Tạo, đọc, cập nhật và xóa; hệ thống ưu tiên ngừng hoạt động để giữ lịch sử'),('Metrics / logs','Số liệu tổng hợp theo thời gian / bản ghi sự kiện chi tiết'),('PromQL / LogQL','Ngôn ngữ truy vấn metrics Prometheus / log Loki'),('CVE / SBOM','Mã lỗ hổng được công bố / danh mục thành phần phần mềm'),('WSL / VPS','Linux trên Windows / máy chủ ảo; VPS không thuộc phạm vi hiện tại')],[3.2,12.8])

chapter('CHƯƠNG 1. TỔNG QUAN KIẾN TRÚC VÀ CHUẨN BỊ MÔI TRƯỜNG')
h('1.1. Giới thiệu đề tài và vấn đề cần giải quyết')
p('Đơn vị tư vấn cần quản lý cuộc hẹn giữa khách hàng và nhiều chuyên viên. Nếu tiếp nhận lịch bằng tin nhắn hoặc một biểu mẫu không có tài khoản, người điều phối khó xác định chủ sở hữu lịch, dễ nhận trùng thời gian và khó theo dõi yêu cầu nào đã được xác nhận. Thông tin dịch vụ, thời lượng và chuyên viên cũng có thể thay đổi trong khi lịch cũ vẫn phải giữ được ý nghĩa ban đầu.')
p('Hệ thống đặt lịch hẹn giải quyết các vấn đề trên bằng quy trình có kiểm soát: admin chuẩn bị dịch vụ và phân công nhân viên; mở các khung giờ hợp lệ; khách đăng nhập để gửi yêu cầu; admin xác nhận và hoàn tất sau khi buổi tư vấn kết thúc. Tài khoản khách giúp tra cứu lịch của chính mình, cập nhật liên hệ và hủy trong thời hạn được phép. Lịch sử thay đổi trạng thái được lưu để giải thích ai đã thao tác và vào thời điểm nào.')
p('Trong đề tài này, một chuyên viên có thể phụ trách nhiều dịch vụ, một dịch vụ có thể được nhiều chuyên viên đảm nhiệm. Hai người có thể nhận khách cùng giờ nhưng một người không được có hai khung giờ đang mở chồng nhau. Đây là khác biệt quan trọng giữa việc chọn một mốc giờ chung cho toàn đơn vị và việc đặt hẹn với nguồn lực cụ thể.')
h('1.2. Yêu cầu của đề 40 và phạm vi thực hiện')
table('Đối chiếu yêu cầu chính và sản phẩm triển khai',['Yêu cầu','Cách thực hiện và điều kiện nghiệm thu'],[
('Website dịch vụ, khung giờ, khách hàng, xác nhận','Frontend + backend; tài khoản khách; quản trị 5 phân hệ; luồng pending → confirmed → completed hoặc cancelled'),('PostgreSQL và pgAdmin','Nghiệp vụ/session lưu DB; migration 1/2/3; pgAdmin truy vấn bảng thật'),('Đóng gói Compose','13 service; 11 chạy dài hạn, migrate/log-init exit 0; healthcheck và volume'),('Nginx và bảo mật web','Reverse proxy, HTTPS tự ký 8443, headers và chặn endpoint nội bộ'),('Prometheus và Grafana','5 job UP; dashboard 15 panel gồm CPU, RAM, Network I/O, web và DB'),('Loki, Promtail và LogQL','Ba nguồn log cốt lõi, truy vấn có kết quả và pipeline thời gian sự kiện'),('Hardening và GitHub','Nhiều lớp bảo vệ; source/config/README và lịch sử commit có ý nghĩa'),('Báo cáo và demo','Bìa cá nhân, mục lục, danh mục hình/bảng, sơ đồ, ảnh thật và kịch bản')],[5.1,10.9])
p('Nguồn yêu cầu chính là đề thực hành số 40 [1]. Bao_Cao_Trien_Khai_He_Thong.docx [2] được dùng để đối chiếu bốn nhóm nội dung: chuẩn bị kiến trúc, triển khai web/DB/proxy, monitoring/logging và hardening. Bao_cao_Appointment_System.docx [3] cung cấp bố cục bìa và cách tổ chức một báo cáo có cơ sở công nghệ, phân tích thiết kế, triển khai, kiểm thử và phụ lục. Các mô tả của mẫu được điều chỉnh theo cấu hình thật; không tự nhận microservices, UFW hay log mọi container khi chưa triển khai.')
p('Mục tiêu vận hành là máy cá nhân, không có VPS/domain công khai. Website có HTTPS tại https://localhost:8443; HTTP 8080 giữ làm kênh lab tương thích với các script. Cổng quản trị công cụ chỉ bind 127.0.0.1. Nhân viên là đối tượng được admin quản lý và phân công, chưa có tài khoản riêng. Email/SMS, xác minh email, quên mật khẩu, thanh toán và đồng bộ lịch ngoài nằm ngoài phiên bản hiện tại.')
h('1.3. Vai trò các công nghệ trong hệ thống')
table('Chức năng và lý do sử dụng từng thành phần',['Thành phần','Chức năng cụ thể trong dự án'],[
('Docker / Compose','Đóng gói tiến trình; mô tả service, network, volume, healthcheck và dependency trong một cấu hình có thể dựng lại'),('Node.js / Express','Nhận HTTP, kiểm tra dữ liệu, xác thực/phân quyền, xử lý transaction và tạo phản hồi; Node 24.22.1 trong image hiện tại'),('EJS / CSS / JavaScript','Render trang từ backend, hiển thị biểu mẫu/danh sách, gọi API và thông báo kết quả; không thay kiểm tra quyền server'),('PostgreSQL / pgAdmin','Lưu dữ liệu nhất quán / cung cấp giao diện quản trị DB; pgAdmin không xử lý nghiệp vụ đặt hẹn'),('Nginx','Cổng web, TLS termination, chuyển tiếp request, headers, rate limit và giấu endpoint nội bộ'),('Prometheus / exporters','Thu metrics định kỳ; exporter chuyển số liệu Nginx/DB sang định dạng Prometheus'),('cAdvisor','Lấy CPU, RAM và lưu lượng mạng theo container trên Docker host'),('Grafana','Truy vấn datasource để hiển thị dashboard metrics và tra cứu log'),('Promtail / Loki','Đọc file log và đẩy về kho log; Loki lưu, lập chỉ mục nhãn và thực thi LogQL'),('Git / GitHub','Lưu phiên bản source và cấu hình; cho phép đối chiếu thay đổi theo commit và dựng lại từ repo')],[4.2,11.8])
p('Một chức năng nghiệp vụ có thể liên quan nhiều thành phần nhưng mỗi thành phần có trách nhiệm khác nhau. Ví dụ, khi khách gửi yêu cầu đặt lịch, Nginx nhận kết nối HTTPS, Express kiểm tra session/CSRF và PostgreSQL quyết định tính nhất quán của transaction. Prometheus chỉ quan sát số request và độ trễ; Loki giúp xem lại sự kiện, không phải nơi lưu lịch hẹn. Tách đúng vai trò giúp tìm lỗi đúng tầng thay vì quy mọi lỗi cho giao diện.')
h('1.4. Kiến trúc tổng thể và kết nối toàn bộ hệ thống')
fig(D/'00-full-system-architecture.png','Kiến trúc tổng thể đủ 13 service web dữ liệu metrics và logs','Sơ đồ tổng thể gồm ba vùng liên kết trong cùng stack. Vùng A xử lý web và dữ liệu; vùng B thu metrics bằng Prometheus và trình bày bằng Grafana; vùng C đọc ba nguồn log bằng Promtail rồi lưu tại Loki. Grafana truy vấn cả Prometheus và Loki. Mũi tên xanh dương là chiều gửi yêu cầu, xanh lá là chiều dữ liệu log; kết quả truy vấn đi ngược chiều yêu cầu. App /metrics ở vùng B là endpoint của app vùng A, không phải container thứ mười bốn. Hai job log-init và migrate chạy hoàn tất rồi dừng; các volume không được tính thành service.',max_height=22)
p('Toàn bộ 13 service trong compose.yaml đều xuất hiện trên sơ đồ: log-init, postgres, migrate, app, nginx, pgadmin, postgres-exporter, nginx-exporter, cadvisor, prometheus, loki, promtail và grafana. Sơ đồ dưới đây phóng lớn riêng luồng web; các luồng metrics và logs được phân tích tiếp ở Chương 3 để giữ chữ và đường nối dễ đọc.')
table('Các kết nối giữa thành phần và ý nghĩa chiều mũi tên',['Bên gửi → bên nhận','Giao thức và đích','Mục đích'],[
('Trình duyệt → Nginx','HTTPS host 8443; HTTP 8080 phụ','Mở website, đăng nhập và gửi thao tác'),
('Nginx → app','HTTP app:3000','Chuyển request sau khi kết thúc TLS'),
('app / migrate / pgAdmin → postgres','SQL postgres:5432','Nghiệp vụ/session; migration; quản trị DB'),
('nginx-exporter → Nginx','HTTP nginx:8081/stub_status','Đọc thống kê proxy rồi xuất /metrics'),
('postgres-exporter → postgres','SQL postgres:5432','Đọc thống kê bằng role appointment_monitor'),
('Prometheus → app','GET app:3000/metrics','Lấy request, latency, Node.js'),
('Prometheus → các exporter','GET :9113/metrics và :9187/metrics','Lấy số liệu Nginx và PostgreSQL'),
('Prometheus → cAdvisor / chính nó','GET cadvisor:8080/metrics; prometheus:9090/metrics','Tài nguyên container / sức khỏe collector'),
('Grafana → Prometheus','HTTP prometheus:9090; PromQL','Lấy chuỗi thời gian cho dashboard'),
('Ba volume log → Promtail','Đọc file trong mount :ro','Theo dõi log hiện hành và nhớ vị trí đọc'),
('Promtail → Loki','POST loki:3100/loki/api/v1/push','Đẩy batch log có nhãn và timestamp'),
('Grafana → Loki','HTTP loki:3100; LogQL','Tra cứu log tại Explore và panel log'),
('Trình duyệt quản trị → các UI','Host 3001 / 9090 / 5050','Grafana / Prometheus / pgAdmin; chỉ loopback')],[5.3,5.3,5.4])
h3('1.4.1. Luồng truy cập web và dữ liệu')
fig(D/'01-web-data-flow.png','Luồng truy cập HTTPS, backend và cơ sở dữ liệu','Mũi tên trong hình là chiều bên chủ động gửi yêu cầu. Trình duyệt gửi HTTPS vào Nginx ở cổng host 8443; Nginx giải mã TLS và proxy HTTP sang app:3000. Backend gửi SQL tới postgres:5432. Quản trị viên truy cập pgAdmin qua cổng host 5050; bên trong container, pgAdmin nghe 8080 và kết nối DB qua mạng db. Phản hồi đi ngược các mũi tên. Cổng host và cổng container được ghi riêng để tránh nhầm 5050 với 8080.')
p('Ứng dụng được tổ chức như một backend nguyên khối có module, render frontend EJS. Mô hình này phù hợp quy mô học phần vì giảm chi phí triển khai giao tiếp giữa các dịch vụ nghiệp vụ và vẫn cho phép kiểm tra transaction tập trung. Các container hạ tầng được triển khai riêng không đồng nghĩa đã tách hệ thống nghiệp vụ thành microservices. Việc gọi đúng kiến trúc là cần thiết để đánh giá đúng phạm vi thực hiện.')
p('Trong source, server.js tổ chức middleware, metric/log và API đặt lịch; customer-auth.js xử lý tài khoản, đăng nhập và lịch của khách; management.js tổ chức các phân hệ quản trị; domain.js chứa validation, mật khẩu và chuyển trạng thái; db.js quản lý kết nối PostgreSQL. Các view và tài nguyên tĩnh thuộc app/views và app/public. Người dùng thao tác qua giao diện nhưng các kiểm tra quyết định vẫn nằm tại backend và DB.')
h('1.5. Mạng Docker, cổng và môi trường chuẩn bị')
table('Phân tách network và cổng truy cập',['Phạm vi','Cấu hình','Ý nghĩa'],[
('frontend','Nginx và app','Nhận request web; app không publish 3000 lên Windows'),('db – internal','DB, app/migrate, pgAdmin, DB exporter','Tên dịch vụ postgres dùng DNS Docker; 5432 không publish'),('monitoring – internal','App/proxy, exporters, cAdvisor, Prometheus, Loki, Promtail, Grafana','Trao đổi metrics/log trong Docker; không mở trực tiếp các exporter/Loki'),('management','pgAdmin, Grafana, Prometheus','Đưa UI công cụ ra loopback Windows'),('HTTPS / HTTP web','127.0.0.1:8443 / :8080','HTTPS là địa chỉ giới thiệu trong báo cáo; HTTP dùng kiểm thử lab'),('UI công cụ','127.0.0.1:5050 / :3001 / :9090','pgAdmin / Grafana / Prometheus; không public LAN/Internet')],[4.0,6.0,6.0])
p('Docker Compose cung cấp phân giải tên dịch vụ và cho phép khai báo các network riêng [4]. Trong cùng container, localhost chỉ trỏ về chính container đó; vì vậy app phải dùng postgres để kết nối DB, không dùng localhost:5432. Các network internal giới hạn kết nối ở phạm vi network nhưng không thay cho mọi cơ chế phân quyền: app vẫn có đường vào DB nên cần role DB và validation phù hợp.')
p('Máy thực hành chạy Windows với WSL 2 và Docker Desktop Linux containers; Engine được ghi nhận 29.9.2 tại lần kiểm tra. Node.js 24 phục vụ các script host, Git/OpenSSL tạo cấu hình và chứng chỉ. Docker Engine phải chạy thật; chỉ có lệnh docker trên PATH chưa đủ. Trước khi dựng cần kiểm tra cổng còn trống, dung lượng image/volume và mạng tải package/image. Không ghi cấu hình RAM/CPU máy khi chưa có số liệu kiểm tra tương ứng.')
h('1.6. Cấu trúc repository và cách dựng lần đầu')
code('app/                 Backend, schema, view, tài nguyên frontend, Dockerfile\ninfra/nginx/         Reverse proxy, TLS local, xoay log\ninfra/prometheus/    Scrape jobs\ninfra/grafana/       Provision datasource và dashboard\ninfra/loki/          Lưu trữ và retention log\ninfra/promtail/      Nguồn log và pipeline\nscripts/ tests/      Setup, vận hành, kiểm thử và nghiệm thu\ndocs/ reports/       Phân tích, minh chứng và báo cáo\ncompose.yaml         Toàn bộ service/network/volume')
code('git clone https://github.com/vzmanhdev/dtc245200223_dat_lich_hen.git\ncd dtc245200223_dat_lich_hen\nnpm.cmd ci --ignore-scripts\nnpm.cmd run setup\npowershell -ExecutionPolicy Bypass -File scripts/start.ps1')
p('setup-env.js sinh secret riêng vào .env và giữ nguyên nếu file đã tồn tại. setup-tls.js tạo cặp chứng chỉ/key localhost khi chưa có, kiểm tra SAN và hạn dùng khi sử dụng lại. Các file mật khẩu, key, dump và tài liệu gốc của học phần bị gitignore. Trong build context, .env, cert/key và tài liệu không cần thiết cũng được loại trừ. README là điểm bắt đầu để người khác clone và dựng lại; mật khẩu của máy hiện tại không được dùng làm cấu hình chung cho mọi người.')

chapter('CHƯƠNG 2. TRIỂN KHAI ỨNG DỤNG WEB CƠ SỞ DỮ LIỆU VÀ HTTPS')
h('2.1. Phân tích chức năng và phân quyền người dùng')
table('Ma trận chức năng theo tác nhân',['Chức năng','Khách công khai','Khách đăng nhập','Admin'],[
('Xem dịch vụ / chuyên viên / lịch trống','Có','Có','Có'),('Đăng ký tài khoản khách','Có','Chuyển về tài khoản','Không tạo quyền admin qua form'),('Đặt lịch','Yêu cầu đăng nhập','Có, lấy hồ sơ từ DB','Điều phối qua quản trị'),('Xem/hủy lịch cá nhân','Không','Chỉ lịch sở hữu; hủy trước giờ bắt đầu','Xem và xử lý theo quyền admin'),('Dịch vụ / nhân viên / phân công','Không','Không','Thêm, sửa, ngừng và bảo vệ lịch đang giữ'),('Mở/đóng khung giờ','Không','Không','Lịch đơn hoặc hàng loạt, chống overlap'),('Khóa/mở tài khoản khách','Không','Không','Có; không tự hủy các lịch đã đặt'),('Xác nhận / hoàn tất lịch','Không','Không','Trước giờ bắt đầu / sau giờ kết thúc')],[5.2,3.2,3.8,3.8])
p('Hệ thống dùng một đường dẫn /login cho cả khách và admin. Email và mật khẩu được xác thực, vai trò được xác định từ DB và session được thay mới sau đăng nhập. Admin được chuyển đến /admin, khách đến /account. Form đăng ký chỉ tạo customer_accounts, không có trường do client gửi để tự cấp quyền quản trị. Cách này giảm nhầm giữa username admin cũ và email của khách.')
p('Middleware tải lại trạng thái active và danh tính từ DB theo session, vì vậy tài khoản bị khóa không tiếp tục hoạt động chỉ dựa vào cookie cũ. Server kiểm tra ownership khi đọc/hủy lịch và kiểm tra role khi thay danh mục hoặc trạng thái. Việc ẩn nút trên frontend chỉ hỗ trợ trải nghiệm; một request tự tạo bằng công cụ ngoài trình duyệt vẫn phải chịu cùng kiểm tra quyền.')
fig(E/'login-unified-current.png','Giao diện đăng nhập thống nhất theo email và vai trò','Biểu mẫu có email và mật khẩu, không yêu cầu khách tự chọn vai trò. Phần hướng dẫn chỉ rõ server mở đúng không gian theo quyền tài khoản; liên kết đăng ký tạo tài khoản khách. Ảnh không phải bằng chứng duy nhất về phân quyền: quyền server được kiểm tra thêm bằng integration test. Các ví dụ thao tác trong phần này có thể dùng HTTP của lab; cấu hình và kiểm tra HTTPS được trình bày riêng ở mục 2.6.')
h('2.2. Quy trình đặt hẹn và quản lý trạng thái')
fig(D/'05-appointment-states.png','Các trạng thái và nhánh chuyển của lịch hẹn','pending được tạo khi khách gửi yêu cầu hợp lệ; confirmed do admin xác nhận; completed chỉ được admin thực hiện sau giờ kết thúc. Khách có thể hủy trước giờ bắt đầu, admin có thể hủy pending/confirmed kể cả quá giờ để xử lý tồn đọng. cancelled giữ lịch sử nhưng không tiếp tục giữ slot; completed cũng là trạng thái kết thúc. Không có thao tác mở lại lịch kết thúc trong luồng hiện tại.')
p('Quy trình thực tế bắt đầu từ danh mục: admin tạo dịch vụ, chọn nhân viên phụ trách rồi mở khung giờ. Khách đăng nhập, chọn dịch vụ, chuyên viên, ngày và slot còn trống. Khi gửi, backend lấy tên, điện thoại và email từ hồ sơ DB thay vì tin nội dung định danh do client cung cấp. Kết quả thành công trả mã lịch; khách theo dõi trạng thái trong tài khoản. Admin lọc theo mã hoặc khách để xác nhận và tiếp tục điều phối buổi tư vấn.')
p('Dữ liệu demo đã thực hiện là lịch MH-75CB5CEF4D4C, Khách Demo, dịch vụ Tư vấn triển khai hệ thống, chuyên viên Phạm Gia Huy, 27/10/2026 09:00–10:00, phí snapshot 300.000 đồng và trạng thái confirmed. Ví dụ này chứng minh khách đặt với chuyên viên cụ thể, không phải đặt một giờ vô danh. Khi demo sau ngày đó phải mở slot tương lai mới; không sửa giờ hệ thống để làm nút xác nhận xuất hiện.')
fig(E/'admin-confirmation-desktop.png','Quản trị lọc mã lịch và hiển thị kết quả xác nhận','Thanh điều hướng có năm phân hệ: lịch hẹn, dịch vụ, nhân viên, khung giờ và khách hàng. Bộ lọc đang tìm một mã demo, nên bảng chỉ hiện lịch cần kiểm tra. Cột dịch vụ/chuyên viên, ngày giờ và trạng thái giúp đối chiếu cùng một cuộc hẹn. Các thẻ tổng hợp vẫn phản ánh toàn DB tại thời điểm chụp, gồm cả dữ liệu kiểm thử; không suy ra chúng chỉ đếm kết quả đã lọc.')
h('2.3. Quản lý dịch vụ nhân viên khung giờ và hồ sơ khách')
p('Dịch vụ chứa tên, mô tả, giá và thời lượng. Thời lượng là bội số 15 phút trong khoảng 15–240 phút; giá được kiểm tra trong giới hạn triển khai. Sau khi thêm dịch vụ, admin phải phân công nhân viên và mở lịch trước khi khách có thể đặt. Đổi thời lượng khi còn slot tương lai đang mở bị từ chối, tránh làm thời lượng slot không còn khớp với dịch vụ.')
p('Nhân viên có họ tên, chức danh, giới thiệu và trạng thái nhận lịch. Bảng staff_services lưu phân công nhiều–nhiều. Bỏ phân công hoặc ngừng dịch vụ/nhân viên khi còn lịch tương lai pending/confirmed bị chặn để bảo vệ cam kết với khách. Những đối tượng đã phát sinh lịch không được xóa cứng bằng chức năng thông thường; admin ngừng và có thể mở lại khi phù hợp.')
fig(E/'admin-staff-current.png','Quản lý chuyên viên và các dịch vụ được phân công','Biểu mẫu phía trên tạo hồ sơ và chọn dịch vụ phụ trách; bảng phía dưới cho thấy từng người được giao dịch vụ nào, đang nhận lịch hay đã ngừng. Các bản ghi gắn nhãn kiểm thử là dữ liệu do bộ kiểm thử tạo và đã ngừng, không phải nhân viên thật. Phân công là điều kiện cần để mở slot cho cặp chuyên viên–dịch vụ.')
p('Khung giờ gắn với đúng nhân viên và dịch vụ, có start_at, end_at và trạng thái mở. Có thể tạo đơn hoặc hàng loạt tối đa 32 slot cùng ngày trong giới hạn 365 ngày. Thao tác hàng loạt chạy trong một transaction; một slot không hợp lệ làm cả lô bị rollback. Slot đã có lịch sử không được tùy ý sửa người, dịch vụ hoặc thời gian; cần xử lý lịch liên quan và tạo slot mới khi điều chỉnh kế hoạch.')
p('Khách hàng có tài khoản riêng và lịch sử được truy vấn theo tài khoản. Admin tìm, sửa tên/điện thoại và khóa/mở tài khoản; khóa không tự hủy các lịch đã đặt, vì hai hành động có ý nghĩa khác nhau. Khách tự chỉnh hồ sơ ở /account. Các snapshot liên hệ và lịch sử vẫn được giữ để phản ánh thông tin tại thời điểm tạo lịch.')
fig(E/'customer-account-confirmed.png','Khách theo dõi lịch hẹn đã xác nhận trong tài khoản','Ảnh mobile cho thấy mã lịch, chuyên viên, thời gian và nhãn đã xác nhận. Nút hủy chỉ dành cho lịch hợp lệ theo quyền và thời điểm; backend kiểm tra lại khi nhận request. Đây là tài khoản demo với email example.com. Hệ thống cập nhật trạng thái khi tải lại hoặc gọi API, chưa gửi email/SMS xác nhận tự động.',width=8.5)
h('2.4. Thiết kế PostgreSQL và chống đặt trùng')
fig(D/'04-actual-erd.png','Sơ đồ quan hệ đúng theo schema PostgreSQL đang chạy','Tài khoản sở hữu lịch theo đường customer_accounts → customers → appointments, không có cột appointments.account_id. appointments tham chiếu customers bằng customer_id và time_slots bằng slot_id. Một slot có thể xuất hiện trong nhiều lịch sử đã hủy nhưng chỉ có một lịch đang giữ theo unique index. staff_services biểu diễn phân công; time_slots tham chiếu riêng staff và services, đồng thời backend kiểm tra cặp phân công hợp lệ.')
table('Các bảng nghiệp vụ và dữ liệu được bảo toàn',['Bảng','Nội dung và quan hệ quan trọng'],[
('customer_accounts','Email, password_hash, hồ sơ và active; là danh tính đăng nhập khách'),('customers','Liên hệ snapshot của lượt đặt, account_id nullable để giữ lịch cũ chưa có tài khoản'),('services / staff','Danh mục dịch vụ và nhân viên; active thay cho xóa cứng'),('staff_services','Khóa ghép staff_id/service_id; phân công nhiều–nhiều'),('time_slots','Cặp nhân viên/dịch vụ, timestamptz đầu/cuối, is_open và constraint thời gian'),('appointments','Mã lịch/token hash/idempotency, customer_id, slot_id, status và snapshot tên/người/giá'),('appointment_events','Trạng thái trước/sau, tác nhân và thời điểm; không ghi mật khẩu'),('admin_users / session / schema_migrations','Tài khoản quản trị / phiên server-side / phiên bản migration')],[4.9,11.1])
p('Để chống tranh chấp, request đặt lịch bắt đầu transaction, khóa theo idempotency key và kiểm tra yêu cầu đã được tạo trước đó hay chưa. Tiếp theo SELECT FOR UPDATE khóa slot, dịch vụ và nhân viên; backend kiểm tra hoạt động, thời gian, thời lượng, phân công và các ID client chọn. Khi hợp lệ mới thêm customers, appointments và sự kiện pending rồi COMMIT. Lỗi giữa chừng dẫn tới ROLLBACK, không để thông tin khách được thêm mà không có lịch tương ứng.')
code('-- Trích điều kiện bảo vệ ở tầng DB\nCREATE UNIQUE INDEX one_held_booking_per_slot ON appointments(slot_id)\nWHERE status IN (\'pending\',\'confirmed\',\'completed\');\n-- V3: không chồng các khoảng đang mở của cùng nhân viên\nEXCLUDE USING gist (staff_id WITH =,\n  tstzrange(start_at,end_at,\'[)\') WITH &&) WHERE (is_open);')
p('Unique index là lớp chặn cuối khi hai request cùng tranh một slot. Exclusion constraint xử lý các khoảng thời gian chồng nhau của cùng nhân viên, nhưng cho phép người khác nhận khách cùng giờ. Idempotency xử lý trường hợp bấm gửi lại vì mạng chậm: cùng key và cùng nội dung trả lịch cũ; dùng lại key cho nội dung khác bị từ chối. Những cơ chế này bổ sung cho transaction/isolation của PostgreSQL [5], không được thay bằng việc chỉ kiểm tra slot trống trên frontend.')
p('Migration 1/2/3 được thiết kế cộng thêm và giữ dữ liệu cũ. Tên dịch vụ, chuyên viên và giá được snapshot khi tạo lịch để thay đổi danh mục không viết lại lịch sử. Lịch cũ chưa có tài khoản không tự gán theo email chưa xác minh. Role app không phải superuser hoặc chủ schema; migration có quyền riêng, monitoring dùng role tách biệt. DB dùng timestamptz, giao diện trình bày giờ Việt Nam UTC+7.')
h('2.5. Đóng gói Docker Compose và công cụ pgAdmin')
p('Dockerfile hiện dùng multi-stage build. Stage đầu dùng Node 24 trên Alpine 3.23 để npm ci theo lockfile, chỉ lấy dependency production và không chạy install script. Stage runtime dùng Alpine cùng nhánh, cập nhật package, cài libstdc++, chép Node binary cùng node_modules và source. npm/Yarn chỉ phục vụ build, không được đưa vào runtime. App chạy UID node 1000, filesystem chỉ đọc và có volume log riêng; đây là bản đã khắc phục 13 cảnh báo High ở image cũ.')
table('Service Compose và vai trò vận hành',['Service','Vai trò','Trạng thái mong đợi'],[
('log-init','Tạo quyền ghi phù hợp cho volume log trước khi dịch vụ chạy','Exited 0'),('postgres','Lưu DB và session; pg_isready làm healthcheck','Running/healthy'),('migrate','Áp dụng schema và khởi tạo có điều kiện','Exited 0'),('app / nginx','Backend/frontend / cổng HTTP-TLS','Running/healthy'),('pgadmin','Công cụ quản trị DB qua giao diện','Running/healthy'),('nginx-exporter / postgres-exporter','Chuyển số liệu dịch vụ sang metrics','Running'),('cadvisor','Quan sát CPU/RAM/network container','Running'),('prometheus / grafana','Thu metrics / dashboard và Explore','Running'),('loki / promtail','Kho log / tác nhân đọc và chuyển log','Running')],[4.5,8.3,3.2])
p('depends_on với điều kiện healthy hoặc service_completed_successfully buộc thứ tự khởi động hợp lý: DB sẵn sàng trước migration, migration hoàn thành trước app, app healthy trước Nginx. Exited 0 của hai job là bình thường; exited của service chạy dài hạn là lỗi. Restart policy không thay cho healthcheck: container có tiến trình nhưng không phục vụ được request vẫn cần bị phát hiện. Bộ verify-stack kiểm tra đủ danh sách service dự kiến thay vì chỉ các container đang xuất hiện.')
fig(E/'pgadmin-services-verified.png','pgAdmin đọc bảng services trong cơ sở dữ liệu thật','Server trong pgAdmin dùng host postgres, database appointment và cổng nội bộ 5432. Query Tool cho thấy bảng services cùng Query complete. Ảnh chụp trước các lượt kiểm thử sau nên số dòng có thể khác hiện tại. Để nghiệm thu có thể truy vấn thêm staff, time_slots, appointments và appointment_events; không chụp bảng password_hash hoặc thông tin riêng của khách thật vào báo cáo.')
code('SELECT id, name, duration_minutes, price, active\nFROM services ORDER BY id;\nSELECT status, count(*) FROM appointments GROUP BY status;')
h('2.6. Cấu hình Nginx reverse proxy và HTTPS tự ký')
p('HTTPS bảo vệ kết nối giữa trình duyệt và Nginx bằng TLS. Nginx là điểm kết thúc TLS, vì vậy chứng chỉ/key nằm ở proxy và backend vẫn dùng HTTP trên mạng Docker nội bộ. Mô hình này không phải mã hóa toàn bộ mọi hop: đoạn Nginx–app và app–DB trong lab chưa dùng TLS riêng. Báo cáo phân biệt rõ nơi mã hóa và nơi dựa vào network isolation để tránh mô tả vượt cấu hình thực tế.')
code('server {\n    listen 8080;\n    listen 8443 ssl;\n    ssl_certificate /etc/nginx/certs/localhost.crt;\n    ssl_certificate_key /etc/nginx/certs/localhost.key;\n    ssl_protocols TLSv1.2 TLSv1.3;\n    resolver 127.0.0.11 valid=5s ipv6=off;\n    set $appointment_backend app:3000;\n    location / {\n        proxy_pass http://$appointment_backend;\n        proxy_set_header Host $http_host;\n        proxy_set_header X-Forwarded-Proto $scheme;\n    }\n}')
note('Đoạn trên là trích rút gọn để giải thích; file infra/nginx/default.conf là cấu hình đầy đủ. SSL/TLS và vị trí certificate/key tương ứng cách cấu hình Nginx chính thức [6].')
p('Cổng HTTPS trên host là 8443 vì lab dùng port cao; HTTPS không bắt buộc phải dùng 443. Server hỗ trợ TLS 1.2/1.3 và tắt thông tin phiên bản Nginx. X-Forwarded-Proto cho backend biết request gốc đi qua HTTPS; Express chỉ tin một hop proxy và cookie session dùng secure:auto, nhờ đó có Secure khi qua TLS. Cặp chứng chỉ được mount chỉ đọc và không đưa private key lên Git. Không dùng password/key cố định cho mọi checkout.')
from cryptography import x509
from cryptography.hazmat.primitives import hashes
certificate=x509.load_pem_x509_certificate((ROOT/'infra/nginx/certs/localhost.crt').read_bytes())
cert={'sha256':certificate.fingerprint(hashes.SHA256()).hex()}
table('Thông tin chứng chỉ HTTPS local được kiểm tra',['Thuộc tính','Giá trị thực tế'],[
('Địa chỉ','https://localhost:8443'),('Subject / Issuer','O=Appointment Local Lab, CN=localhost / cùng chủ thể tự ký'),('SAN','DNS:localhost; IP Address:127.0.0.1'),('Khóa / hiệu lực','RSA 2048; từ 24/10/2026 06:56:59 đến 24/10/2027 06:56:59 (UTC+7)'),('SHA256 fingerprint',' '.join(cert['sha256'][i:i+8] for i in range(0,64,8))),('Kiểm tra TLS','HTTP 200; authorized=true; TLSv1.3; cookie Secure; /metrics public trả 404')],[4.7,11.3])
p('Chứng chỉ tự ký vẫn tạo kết nối mã hóa nhưng chưa được trình duyệt tin cậy như CA công cộng. Vì vậy trình duyệt có thể hiển thị cảnh báo trước khi truy cập. Người vận hành phải kiểm tra đúng localhost và fingerprint trước khi tự quyết định tin chứng chỉ; báo cáo không coi thao tác bỏ cảnh báo là bằng chứng đã có CA tin cậy. Không tự cài CA vào Windows, không tắt xác minh TLS và không dùng NODE_TLS_REJECT_UNAUTHORIZED=0.')
p('verify-tls.js kiểm tra trực tiếp Nginx từ mạng Docker bằng đúng certificate CA cục bộ, servername localhost và rejectUnauthorized=true. Cách này tránh trường hợp antivirus trên Windows thay certificate của kết nối localhost. Kết quả xác nhận TLS 1.3, hostname hợp lệ, cookie Secure và endpoint metrics bị chặn; không suy ra mọi trình duyệt trên host đã tin CA. HTTP 8080 vẫn được giữ cho một số script lab, chưa ép chuyển hướng toàn bộ sang HTTPS.')
code('npm.cmd run setup\ndocker compose exec -T nginx nginx -t\nnode scripts/verify-tls.js\n# Kết quả thực tại lần nghiệm thu:\n# PASS TLS certificate verification, Secure cookie, private metrics blocked TLSv1.3')
https_image=E/'https-browser-verified.png'
if https_image.exists():fig(https_image,'Giao diện website truy cập qua HTTPS localhost','Nội dung trang được chụp tại https://localhost:8443 sau khi phiên trình duyệt đã xử lý việc tin chứng chỉ local. Trang chủ và liên kết đặt lịch được tải qua HTTPS. Ảnh nội dung không bao gồm thanh địa chỉ của trình duyệt; URL phiên được ghi trong minh chứng máy. Kết quả verify-tls bổ sung kiểm tra CA/SAN, giao thức TLS và cookie, không đồng nghĩa chứng chỉ đã được CA công cộng cấp.',width=10)
h('2.7. Security headers giới hạn request và khả năng phục hồi proxy')
table('Tác dụng các cơ chế bảo vệ tại proxy và web',['Cơ chế','Tác dụng','Giới hạn'],[
('nosniff','Ngăn trình duyệt tự đoán MIME khác Content-Type','Cần trả Content-Type đúng'),('SAMEORIGIN','Giảm clickjacking bằng hạn chế nhúng trang trong frame','Không thay phân quyền server'),('CSP','Giới hạn nguồn script/tài nguyên phù hợp giao diện','Không thay escape dữ liệu hoặc vá XSS'),('Referrer-Policy','Hạn chế thông tin referrer gửi sang ngữ cảnh khác','Không ẩn dữ liệu đã đưa vào URL'),('Permissions-Policy','Tắt camera, microphone, geolocation không cần dùng','Không phải firewall hoặc cơ chế xác thực'),('Rate/body/timeout','5 request/s tại Nginx, burst 40; body 16 KiB; timeout proxy','Có thể từ chối cả test dồn dập; chưa là chứng minh chống DDoS'),('Chặn endpoint nội bộ','/metrics, /health/*, /stub_status trả 404 ở cổng public','Exporter/health dùng cổng và mạng nội bộ')],[3.7,6.3,6.0])
p('Không bật lại X-XSS-Protection: 1; mode=block chỉ để giống báo cáo mẫu. Header này đã lỗi thời; ứng dụng dùng CSP và escape EJS, phản hồi hiện có X-XSS-Protection=0 [7]. HSTS không bật cho localhost để tránh ép các ứng dụng khác trên cùng hostname sang HTTPS. Những lựa chọn này được giải thích theo bối cảnh lab, không nhận là cấu hình chuẩn cho mọi hệ thống Internet.')
p('Nginx dùng Docker DNS resolver 127.0.0.11 và tên app:3000 trong biến upstream. Khi app được recreate và đổi IP, proxy cập nhật phân giải thay vì giữ IP cũ. Thử recreate đã ghi nhận web phục hồi mà không reload Nginx. Helper xoay log nằm cùng container Nginx để không phụ thuộc PID namespace của một sidecar có vòng đời khác; chi tiết lỗi và cách thử lại nằm ở chương kiểm thử.')

chapter('CHƯƠNG 3. GIÁM SÁT METRICS VÀ LOG TẬP TRUNG')
h('3.1. Phân biệt mục đích của metrics và log')
p('Metrics là các đại lượng số theo thời gian, phù hợp để trả lời hệ thống có đang hoạt động, tải tài nguyên ra sao và độ trễ có thay đổi hay không. Log là bản ghi từng sự kiện, phù hợp để tìm request, mã lỗi hoặc nguyên nhân của một thao tác. Hai nguồn cần được dùng cùng nhau: dashboard chỉ ra thời điểm bất thường, sau đó LogQL giúp đọc sự kiện tương ứng. Không dùng một ảnh dashboard để chứng minh toàn bộ luồng nghiệp vụ đúng.')
p('Hệ thống không ghi mật khẩu, token hoặc body request vào log HTTP. Label metrics dùng route mẫu, method và status, tránh đưa ID khách hoặc query string làm label có số lượng giá trị tăng không giới hạn. App bỏ qua log/metric HTTP của các endpoint health/metrics để không làm request kiểm tra nội bộ lấn át các request nghiệp vụ. Các giá trị trên dashboard là dữ liệu lab, có thể bị ảnh hưởng bởi integration test và không đại diện lưu lượng khách thật.')
h('3.2. Prometheus exporters và cAdvisor')
fig(D/'02-metrics-pull-flow.png','Chiều yêu cầu scrape metrics và truy vấn Grafana','Mũi tên chỉ chiều yêu cầu: Grafana gửi PromQL đến Prometheus; Prometheus chủ động gửi HTTP scrape đến app, Nginx exporter, PostgreSQL exporter và cAdvisor. Metrics trả về chiều ngược lại. Prometheus còn tự scrape chính nó nên có tổng cộng 5 job. Không vẽ mũi tên Nginx đẩy trực tiếp metrics vào Grafana vì cấu hình không hoạt động theo cách đó.')
p('Prometheus scrape mỗi 15 giây và đánh giá rule theo cùng chu kỳ [8]. Job appointment đọc app:3000/metrics; nginx đọc nginx-exporter:9113; postgres đọc postgres-exporter:9187; cadvisor đọc cadvisor:8080; prometheus đọc chính prometheus:9090. Prometheus lưu chuỗi thời gian vào volume riêng, giới hạn retention 7 ngày và dung lượng 1 GB trong cấu hình hiện tại. Việc target UP chỉ chứng minh scrape thành công, không thay kiểm tra các chức năng ứng dụng.')
code('global:\n  scrape_interval: 15s\n  evaluation_interval: 15s\nscrape_configs:\n  - job_name: appointment\n    static_configs:\n      - targets: ["app:3000"]\n  - job_name: nginx\n    static_configs:\n      - targets: ["nginx-exporter:9113"]\n  # Các job postgres, cadvisor, prometheus ở file đầy đủ.')
p('Nginx exporter chuyển stub_status từ cổng nội bộ 8081 thành các metric như nginx_up, số kết nối và tổng request. PostgreSQL exporter kết nối DB bằng role monitoring để cung cấp pg_up, kết nối và giao dịch. cAdvisor đọc tài nguyên từ Docker host để lấy CPU, memory working set và network bytes theo label service. cAdvisor là ngoại lệ privileged; quyền này thuộc công cụ quan sát, không được cấp cho app.')
h('3.3. Dashboard Grafana và cách đọc các chỉ số')
fig(E/'grafana-dashboard-verified.png','Dashboard có dữ liệu về target web và tài nguyên container','Thẻ Scrape targets UP bằng 5 tương ứng năm job được cấu hình; PostgreSQL connected và Nginx connected bằng 1 là metric exporter. Các đồ thị phía dưới biểu diễn request và CPU/RAM theo thời gian. Kết quả kiểm tra sâu xác nhận các truy vấn Prometheus của dashboard có mẫu hữu hạn; ảnh minh chứng không được dùng để giả định các panel ở ngoài vùng nhìn đều đã đạt.')
table('Nhóm panel dashboard và ý nghĩa vận hành',['Nhóm','Chỉ số hoặc dữ liệu','Cách diễn giải'],[
('Tình trạng scrape','sum(up), pg_up, nginx_up','Đếm target scrape thành công và tình trạng exporter tới dịch vụ'),('Web app','Request/s theo status, histogram p95','Quan sát lượng request và độ trễ; p95 không có dữ liệu khi chưa có mẫu phù hợp'),('Nginx','Active connections và request/s','Quan sát kết nối/tải tại proxy; không đồng nhất hoàn toàn với request app'),('PostgreSQL','Connections và committed transactions/s','Biết DB có kết nối và giao dịch; không nhận đây là toàn bộ phân tích query chậm'),('Container','CPU rate, memory working set','Xem phân bố tài nguyên theo label Compose service'),('Network I/O','Receive/transmit bytes/s, bỏ interface lo','Tốc độ mạng, không phải tổng byte từ khi container khởi động'),('Log app','Panel Loki với service=app','Đọc sự kiện ngay trên dashboard, khác với chuỗi số Prometheus')],[3.8,5.0,7.2])
fig(E/'grafana-network-verified.png','Network receive transmit và log app trên Grafana','Hai đồ thị phía trên đo byte nhận/gửi mỗi giây theo Compose service; trục hiển thị B/s hoặc kB/s. Đường tăng tại thời điểm kiểm thử là hoạt động lab, không tự suy ra nghẽn mạng. Panel phía dưới là log thật được đọc từ Loki, cho thấy dashboard có hai datasource với mục đích khác nhau.')
p('Dashboard appointment-local có 15 panel: 14 panel Prometheus và một panel Loki; datasource và dashboard được provision từ file trong Git. Vì vậy checkout mới không cần cấu hình thủ công từng panel. Tuy nhiên provisioning thành công chưa chứng minh có dữ liệu: bộ audit truy vấn từng biểu thức trong 30 phút, kiểm tra mẫu số hữu hạn và fail nếu panel không có mẫu. Nếu vừa khởi động cần ít nhất hai chu kỳ scrape để rate hoạt động.')
h('3.4. PromQL tiêu biểu và điều kiện kiểm chứng')
code('sum(up)\n\nsum by (status) (rate(http_requests_total[2m]))\n\nhistogram_quantile(0.95,\n  sum by (le) (rate(http_request_duration_seconds_bucket[5m])))\n\nsum by (container_label_com_docker_compose_service) (\n  rate(container_network_receive_bytes_total{\n    container_label_com_docker_compose_service!="",interface!="lo"}[2m]))')
p('sum(up) giúp kiểm tra số target đang scrape thành công; cần đối chiếu với số job dự kiến thay vì coi mọi giá trị dương là đủ. rate tính mức thay đổi của counter trong cửa sổ thời gian, phù hợp request/s hoặc byte/s. histogram_quantile tổng hợp bucket để ước lượng p95; chỉ có ý nghĩa khi có request quan sát được. Không thay thế các giá trị thiếu bằng 0 để làm dashboard trông có dữ liệu.')
p('Khi chạy một bản clone độc lập trên cùng Docker host, cAdvisor có thể thấy cả container bản chính và bản thử. Nếu chỉ group theo service, hai project có thể được cộng chung. Vì vậy cần dừng môi trường thử sau kiểm chứng hoặc lọc thêm project khi phân tích tải riêng. Dashboard hiện được nghiệm thu với bản chính sau khi dừng bản thử; không dùng số liệu giai đoạn hai stack cùng chạy để kết luận nhu cầu tài nguyên riêng của một stack.')
h('3.5. Loki Promtail và pipeline log thực tế')
fig(D/'03-log-collection-flow.png','Ba nguồn log các bước xử lý Promtail và truy vấn Loki','Ba volume log chỉ đọc đưa file hiện hành tới Promtail; Promtail đọc, gắn label và đẩy batch tới Loki bằng HTTP. Grafana Explore gửi LogQL đến Loki để lấy kết quả. Không có đường app trực tiếp ghi metrics vào Loki, và pipeline này không thu stdout của mọi container. Vị trí đã đọc được lưu vào volume positions để hạn chế đọc lại khi restart.')
p('App dùng Pino JSON, Nginx dùng access JSON; PostgreSQL ghi log file theo ngày. Hai nguồn JSON được parse trường time theo RFC3339Nano làm timestamp. PostgreSQL plaintext hiện dùng thời điểm thu. Nhãn service nhận app, nginx hoặc postgres; environment=local phân biệt môi trường. Loki lưu filesystem/TSDB schema v13, retention 168 giờ, giới hạn ingestion rate/burst và lookback 168 giờ; cấu hình không mở cổng Loki ra Windows.')
code('pipeline_stages:\n  - json:\n      expressions:\n        event_time: time\n  - timestamp:\n      source: event_time\n      format: RFC3339Nano\n      action_on_failure: skip\n# App chỉ đọc file đang ghi:\n# __path__: /var/log/appointment/app.log')
p('Chỉ đọc app.log/access.log hiện hành để không đưa file .1/.2 trở lại như nguồn mới. Khi xoay log, file cũ vẫn được giữ theo giới hạn và ứng dụng/proxy tiếp tục ghi file mới. Vị trí đọc cùng timestamp sự kiện là hai yếu tố khác nhau: sửa được vị trí mà chưa parse time vẫn có thể đưa sự kiện cũ vào khoảng thời gian mới. Thử nghiệm đã đối chiếu request_id trước/sau rotation để xác nhận mỗi request mới chỉ có một bản.')
p('Promtail được giữ theo yêu cầu bài thực hành, nhưng công cụ đã kết thúc vòng đời hỗ trợ theo tài liệu Grafana [9]. Vì vậy báo cáo không coi pipeline hiện tại là lựa chọn mặc định cho vận hành Internet dài hạn; hướng phát triển cần thay bằng collector được duy trì. Không tự đổi collector trong lúc nghiệm thu nếu điều đó làm khác yêu cầu đề mà chưa đối chiếu lại.')
h('3.6. Ba truy vấn LogQL và cách đọc kết quả')
h3('3.6.1. Xem log của backend')
code('{service="app", environment="local"}')
p('Truy vấn chọn các stream có đúng label app và local. Kết quả chứa thời gian, level số, request_id, method, route mẫu, status và duration_ms. Có thể đối chiếu một request cụ thể bằng request_id; không dùng email, mật khẩu hay token làm từ khóa lấy từ log vì các trường đó không được ghi vào log HTTP. Pino level 30 là info, 50 là error; không dùng biểu thức tìm chữ ERROR để thay cho level số khi định dạng nguồn khác nhau.')
h3('3.6.2. Lọc HTTP lỗi tại reverse proxy')
code('{service="nginx", environment="local"}\n  | json | __error__="" | status >= 400')
p('Pipeline parse JSON rồi loại lỗi parse bằng __error__="", sau đó lọc trường status. Trong lab có các request 404 cố ý từ smoke và LogQL check. Vì vậy thấy log lỗi không tự chứng minh website hỏng; cần đọc method, thời điểm và ngữ cảnh kiểm thử. Nếu cần chỉ xem server error có thể đổi ngưỡng sang status >= 500, nhưng query đó có thể không có kết quả khi hệ thống không phát sinh 5xx.')
fig(E/'loki-logql-verified.png','Grafana Explore thực thi LogQL lọc HTTP lỗi','Vùng query dùng label service/environment đúng cấu hình và JSON filter theo status. Danh sách phía dưới có log thực tế trả về, không phải phần ví dụ tĩnh của tài liệu. Mỗi dòng có timestamp và nội dung JSON. Ảnh ở thời điểm chụp vẫn có thể chứa log lịch sử trước mốc sửa rotation; kiểm thử không-trùng chỉ được kết luận cho sự kiện mới sau sửa.')
h3('3.6.3. Đếm số sự kiện HTTP lỗi trong năm phút')
code('sum(count_over_time(\n  {service="nginx", environment="local"}\n  | json | __error__="" | status >= 400 [5m]))')
p('count_over_time đếm các bản ghi trong cửa sổ năm phút, sum cộng các stream. Đây là metric query từ log, không phải một counter Prometheus độc lập. Có thể mở một URL không tồn tại để sinh 404, đợi collector chuyển dữ liệu rồi truy vấn. Bộ verify-stack tiêu thụ body phản hồi và retry ingestion có giới hạn, cập nhật thời gian kết thúc query ở mỗi lần thử để không kiểm tra một cửa sổ cũ.')
p('Các log trùng lịch sử trước khi sửa không bị xóa hoặc viết lại; chúng hết theo retention. Vì thế không dùng số đếm lịch sử chưa phân loại để kết luận SLA. Promtail/Loki có thể ready nhưng chưa chuyển xong log mới, nên cần phân biệt readiness của dịch vụ và điều kiện dữ liệu đã xuất hiện trong query. Đây là lý do nghiệm thu kiểm tra kết quả truy vấn chứ không chỉ HTTP 200 của Loki.')

h('3.7. Cấu hình và trách nhiệm từng dịch vụ quan sát')
h3('3.7.1. Prometheus và năm scrape job')
p('Prometheus là kho chuỗi thời gian của stack. Service dùng prom/prometheus:v3.5.0 và file infra/prometheus/prometheus.yml được mount chỉ đọc vào /etc/prometheus/prometheus.yml. Bộ lập lịch scrape gửi GET tới /metrics của từng target; endpoint trả văn bản số liệu theo định dạng Prometheus. Job là tên nhóm thu thập, instance là địa chỉ đích; up=1 thể hiện lần scrape thành công, không khẳng định nghiệp vụ đặt lịch đã đúng. UI trên Windows là http://localhost:9090; trong Docker, Grafana dùng http://prometheus:9090.')
import re
scrape={'scrape_configs':[{'job_name':name,'static_configs':[{'targets':[target]}]} for name,target in re.findall(r'job_name:\s*(\S+)\s+static_configs:\s*- targets:\s*\["([^"]+)"\]',(ROOT/'infra/prometheus/prometheus.yml').read_text(encoding='utf8'))]}
assert len(scrape['scrape_configs'])==5, 'Review scrape config format before rebuilding report'
table('Năm job Prometheus và dữ liệu được thu',['Job','Địa chỉ target','Dữ liệu và kiểm tra'],[(x['job_name'],x['static_configs'][0]['targets'][0],{'prometheus':'Metrics của chính Prometheus; phát hiện vấn đề collector','appointment':'HTTP request, histogram latency và Node.js tại app','nginx':'Metrics do nginx-exporter chuyển từ stub_status','postgres':'Thống kê DB qua postgres-exporter và role monitor','cadvisor':'CPU, memory working set và network của container'}[x['job_name']]) for x in scrape['scrape_configs']],[3.0,5.5,7.5])
p('scrape_interval=15s quyết định tần suất lấy mẫu; evaluation_interval=15s là chu kỳ đánh giá rule nếu có rule được cấu hình. File hiện không khai báo rule_files hoặc Alertmanager nên báo cáo không nhận đã có cảnh báo tự động. TSDB ghi vào /prometheus trên prometheus_data; giới hạn 7d và 1GB đồng thời được áp dụng, dữ liệu có thể bị dọn sớm hơn bảy ngày khi chạm giới hạn dung lượng. Mem_limit 512m là trần cấu hình container, không phải mức RAM thực đo.')
h3('3.7.2. Exporter và cAdvisor')
p('nginx-exporter:1.5.3 kết nối http://nginx:8081/stub_status rồi phục vụ /metrics ở cổng 9113. Stub status là nguồn số liệu của Nginx; exporter là bộ chuyển đổi để Prometheus hiểu được. Vì vậy đường kết nối gồm Prometheus → nginx-exporter → Nginx. Nếu up{job="nginx"}=1 mà nginx_up=0, Prometheus vẫn gọi được exporter nhưng exporter không lấy được dữ liệu từ proxy. Cổng 8081 không được publish ra Windows.')
p('postgres-exporter:v0.17.1 dùng DATA_SOURCE_NAME trỏ postgres:5432/appointment với role appointment_monitor, mật khẩu đọc từ môi trường cục bộ. Exporter nghe 9187 và nằm trên hai mạng db/monitoring, làm cầu đọc thống kê có phân quyền, không thay API nghiệp vụ. pg_up=0 thường cần kiểm tra quyền role, mật khẩu và trạng thái DB; không đưa nguyên connection string có mật khẩu vào báo cáo hoặc màn hình demo.')
p('cAdvisor:0.55.1 chạy --docker_only=true, --housekeeping_interval=15s và --store_container_labels=true. Mount host /, /var/run, /sys, /var/lib/docker và /dev/disk phục vụ quan sát container; các mount khai báo :ro. cAdvisor dùng privileged nên vẫn có quyền rộng: mount chỉ đọc không biến công cụ này thành tiến trình ít đặc quyền như app. Nó phục vụ /metrics ở cadvisor:8080, không có cổng Windows công khai. Nhãn Compose service giúp dashboard cộng số liệu theo loại dịch vụ; nếu nhiều project cùng chạy cần lọc thêm project để tránh cộng chung.')
h3('3.7.3. Grafana và provisioning hai datasource')
p('Grafana:12.1.1 là giao diện truy vấn và hiển thị, không tự thu request từ app. Browser mở http://localhost:3001; port mapping chuyển tới grafana:3000. Datasource được provision từ infra/grafana/provisioning/datasources/datasources.yaml: Prometheus UID prometheus trỏ http://prometheus:9090 và là mặc định; Loki UID loki trỏ http://loki:3100. access:proxy nghĩa là Grafana server truy vấn datasource trong Docker; browser không cần truy cập trực tiếp cổng Loki. Hai datasource editable:false để cấu hình trong Git là nguồn quản lý chính.')
code('Prometheus datasource: uid=prometheus; type=prometheus\n  url=http://prometheus:9090; access=proxy; isDefault=true\nLoki datasource: uid=loki; type=loki\n  url=http://loki:3100; access=proxy\nDashboard provider: folder=Appointment; updateIntervalSeconds=15\n  path=/var/lib/grafana/dashboards\nDashboard: uid=appointment-local; refresh=15s; range=now-30m..now')
p('Dashboard được provision từ infra/grafana/dashboards/appointment.json qua providers trong infra/grafana/provisioning/dashboards/dashboards.yaml. Grafana đọc thư mục mỗi 15 giây; dashboard refresh 15 giây là tần suất trình duyệt làm mới, khác scrape_interval của Prometheus. grafana_data giữ DB nội bộ, tài khoản và trạng thái Grafana; file JSON trong Git giữ thiết kế dashboard. Đăng ký tự do và truy cập anonymous đều tắt; mật khẩu admin lấy từ .env, không ghi trong báo cáo.')
h3('3.7.4. Loki và cơ chế lưu log')
p('Loki:3.5.1 nghe loki:3100, nhận batch từ /loki/api/v1/push và phục vụ truy vấn LogQL. File infra/loki/config.yaml dùng TSDB schema v13 từ mốc 2024-01-01, index theo chu kỳ 24h, object_store filesystem và replication_factor=1. Dữ liệu chunk/index/compactor nằm dưới /loki trên loki_data. Đây là một node local, không có nhân bản hoặc object storage ngoài máy; không coi là cấu hình sẵn sàng cao.')
p('retention_period và max_query_lookback đều 168h. Compactor bật retention_enabled và delete_request_store=filesystem để dọn dữ liệu theo chính sách. ingestion_rate_mb=4 và ingestion_burst_size_mb=8 giới hạn lượng log nhận, không phải số byte đo được của ứng dụng. auth_enabled=false chỉ phù hợp phạm vi internal của lab; Loki không publish port host. Nếu mở ra Internet phải bổ sung lớp xác thực và bảo vệ truy cập trước khi thay phạm vi triển khai.')
h3('3.7.5. Promtail và các bước xử lý log')
p('Promtail:3.5.1 là tác nhân đọc file, khác với Loki là kho log. Config ở infra/promtail/config.yaml; HTTP nội bộ 9080, gRPC tắt bằng grpc_listen_port=0, không publish ra host. Client gửi đến http://loki:3100/loki/api/v1/push. Ba mount app_logs/nginx_logs/postgres_logs đều chỉ đọc, promtail_positions ghi /positions/positions.yaml để nhớ offset. Root + DAC_READ_SEARCH giúp đọc thư mục log do UID ứng dụng khác sở hữu; read_only và cap_drop ALL với capability ngoại lệ được giữ rõ trong Compose.')
table('Ba nguồn log Promtail và ý nghĩa trường dữ liệu',['Job và nhãn','Đường dẫn trong Promtail','Xử lý'],[
('appointment; service=app; environment=local','/var/log/appointment/app.log','Parse JSON time → event_time → timestamp RFC3339Nano; giữ log HTTP/backend'),
('nginx; service=nginx; environment=local','/var/log/appointment-nginx/access.log','Parse JSON time tương tự; status, method và đường dẫn proxy phục vụ truy vấn'),
('postgres; service=postgres; environment=local','/var/log/appointment-postgres/*.log','Đọc log text DB, dùng thời điểm thu hiện tại; chưa có parser timestamp DB riêng')],[4.8,5.8,5.4])
p('Nhãn service/environment giúp chọn stream; các trường JSON như status/request_id được parse khi query và không phải mọi trường đều là label. Với app/Nginx, action_on_failure:skip ở timestamp stage nghĩa là bỏ việc gán timestamp khi không parse được và dùng timestamp sẵn có, không phải tự động loại bỏ cả dòng log. Bảo toàn offset, đọc active file và parse đúng thời gian là ba điều riêng cần kiểm tra khi xoay log. Promtail được giữ theo yêu cầu học phần; tình trạng kết thúc hỗ trợ đã được nêu ở mục 3.5 và hướng thay collector thuộc giai đoạn phát triển sau.')
h('3.8. Danh mục đầy đủ mười lăm panel và truy vấn gốc')
p('Danh mục sau được đọc trực tiếp từ appointment.json khi tạo báo cáo, giữ ID, tiêu đề, datasource và biểu thức thật. Đây là cấu hình truy vấn; kết quả pass tại mốc nghiệm thu nằm trong audit, không được suy ra chỉ từ việc có JSON. Các biểu thức rate cần ít nhất hai mẫu; p95 cần dữ liệu request trong cửa sổ năm phút.')
panel_notes={1:'Số target scrape được; kỳ vọng 5 khi đủ stack.',2:'Kết nối exporter đến PostgreSQL; 1 là thành công.',3:'Exporter đọc được Nginx; 1 là thành công.',4:'Tổng tốc độ request app, đơn vị request/giây.',5:'Tốc độ request tách theo HTTP status để đọc lỗi 4xx/5xx.',6:'Ước lượng độ trễ p95 theo giây; không phải trung bình.',7:'Số kết nối Nginx hiện đang hoạt động.',8:'CPU theo service; 100% tương ứng khoảng một CPU core, không phải luôn toàn host.',9:'Memory working set theo service, đơn vị byte.',10:'Số kết nối DB appointment phân theo datname.',11:'Tốc độ transaction commit mỗi giây, không đồng nhất số lịch đặt.',12:'Request/s tại Nginx; phạm vi khác counter request của app.',13:'Log stream app từ Loki, giữ thời gian/sự kiện để điều tra.',14:'Byte nhận mỗi giây, bỏ loopback và cộng theo service.',15:'Byte gửi mỗi giây, bỏ loopback và cộng theo service.'}
dashboard=json.loads((ROOT/'infra/grafana/dashboards/appointment.json').read_text(encoding='utf8'))
for panel in dashboard['panels']:
    h3(f"Panel {panel['id']} {panel['title']}")
    p('Datasource: '+panel['datasource']['uid']+'. '+panel_notes[panel['id']])
    for target in panel.get('targets',[]):
        # Wrap only at expression syntax boundaries; retain original query meaning.
        expr=target.get('expr','').replace(') (',')\n(').replace(') *',')\n*')
        expr=expr.replace('container_label_com_docker_compose_service!="",','container_label_com_docker_compose_service!="",\n  ')
        code(expr)
h('3.9. Volume cấu hình và dữ liệu bền vững của toàn bộ stack')
table('Chín volume và trách nhiệm lưu trữ',['Volume','Service sử dụng','Dữ liệu và vòng đời'],[
('postgres_data','postgres','DB nghiệp vụ, session, tài khoản và schema; cần backup pg_dump'),
('postgres_logs','postgres ghi; log-init cấp quyền; Promtail đọc','Log DB xoay theo ngày/dung lượng; không thay backup DB'),
('app_logs','app ghi; log-init cấp quyền; Promtail đọc','app.log hiện hành và bản xoay; không ghi mật khẩu/token'),
('nginx_logs','Nginx ghi; log-init cấp quyền; Promtail đọc','Access JSON và log proxy; helper xoay trong container Nginx'),
('pgadmin_data','pgadmin','Trạng thái và tài khoản nội bộ pgAdmin'),
('prometheus_data','prometheus','TSDB metrics; retention 7d / 1GB, không lưu lịch hẹn'),
('grafana_data','grafana','DB/tài khoản Grafana; dashboard/datasource còn có provisioning trong Git'),
('loki_data','loki','Chunk/index/compactor; retention 168h, không tự sao lưu ngoài máy'),
('promtail_positions','promtail','Offset file đã đọc; giữ qua recreate để hạn chế đọc lại')],[4.3,5.1,6.6])
p('Các file cấu hình trong infra được bind mount chỉ đọc; volume dữ liệu là named volume do Compose quản lý. Thay container không đồng nghĩa mất volume, nhưng xóa volume hoặc hỏng ổ đĩa vẫn mất dữ liệu. Đổi COMPOSE_PROJECT_NAME tạo bộ volume tên khác nên có thể thấy hệ thống trống nếu dùng sai project. Backup/restore DB ở Chương 4 chỉ nghiệm thu dữ liệu PostgreSQL; chưa nhận đã có cơ chế backup riêng toàn bộ TSDB/Grafana/Loki.')
h('3.10. Cách truy cập kiểm tra và phân biệt lỗi từng tầng')
table('Đường dẫn vận hành và xử lý lỗi quan sát',['Hiện tượng hoặc thao tác','Nơi kiểm tra','Cách diễn giải và bước xử lý'],[
('Mở Grafana','http://localhost:3001; folder Appointment','Dùng tài khoản trong .env cục bộ; mở dashboard appointment-local hoặc Explore'),
('Mở Prometheus','http://localhost:9090/targets','Đối chiếu 5 target; đọc Last scrape/Error trước khi sửa dashboard'),
('Target DOWN','Trang Targets và Docker service logs','Kiểm tra DNS service/port, container và endpoint /metrics'),
('Target UP nhưng pg_up/nginx_up=0','Exporter → DB / stub_status','Kiểm tra nguồn phía sau exporter và quyền DB; không chỉ nhìn up'),
('Grafana No data','Datasource UID, time range, query trong Prometheus','Chọn đúng 30 phút/UTC+7, chờ ít nhất hai scrape; không điền 0 giả'),
('Explore không có log','service/environment, time range, Promtail/Loki','Phát request hợp lệ, chờ batch, kiểm tra mount active file/positions/timestamp'),
('Loki ready nhưng query chưa thấy sự kiện','Ingestion và kết quả LogQL','Readiness không đảm bảo sự kiện đã index; retry có giới hạn và kiểm tra nội dung'),
('Container app khởi động lại','DB, proxy DNS, metric và log','Đối chiếu trạng thái service, thử web và query sau recreate; không xóa volume để sửa')],[5.1,4.9,6.0])
code('docker compose ps -a\ndocker compose logs --tail 50 prometheus grafana loki promtail\ndocker compose logs --tail 50 nginx-exporter postgres-exporter cadvisor\nnode scripts/verify-stack.js\nnode scripts/audit-requirements.js')
p('Các lệnh đọc log không thay kiểm thử nghiệp vụ và người vận hành cần che thông tin nhạy cảm trước khi chia sẻ. verify-stack thực hiện 16 mục, audit kiểm tra sâu mẫu của dashboard và ba nguồn log; kết quả phải đọc theo thời điểm chạy. UI công cụ dùng HTTP loopback riêng trong lab; HTTPS 8443 hiện bảo vệ website qua Nginx, chưa bao phủ Grafana/Prometheus/pgAdmin. Đây là ranh giới triển khai cần nêu rõ khi thuyết trình.')

chapter('CHƯƠNG 4. HARDENING BẢO MẬT VÀ VẬN HÀNH')
h('4.1. Nguyên tắc áp dụng hardening')
p('Hardening của đề tài được tổ chức theo nhiều lớp: giới hạn nơi lắng nghe, quyền tiến trình, khả năng ghi filesystem, quyền DB và kiểm tra thao tác ứng dụng. Một lớp bị vượt qua không được mặc nhiên cho phép truy cập mọi thành phần còn lại. Tuy vậy, các biện pháp này chỉ giảm rủi ro; không khẳng định non-root hoặc read-only ngăn tuyệt đối container escape, đặc biệt khi có công cụ quan sát dùng privileged.')
table('Ma trận hardening tác dụng và cách kiểm tra',['Lớp','Cấu hình / tác dụng','Minh chứng và giới hạn'],[
('Tiến trình','App UID node 1000; cap_drop ALL; no-new-privileges','docker exec id; inspect. Không bảo đảm ngăn mọi leo thang'),('Filesystem','App/Nginx read_only; tmpfs /tmp; volume log cần ghi','Inspect và runtime. DB vẫn cần volume ghi'),('Mạng','db/monitoring internal; cổng host loopback; không publish 5432','Compose/inspect. Không thay mọi chính sách firewall'),('DB','App không superuser/owner; role monitor tách; migrate riêng','Query quyền. Giữ mật khẩu DB ngoài Git'),('Web','HTTPS, CSP/headers, session HttpOnly/SameSite, Secure qua TLS','verify-tls và response header; CA local chưa mặc nhiên tin cậy'),('Thao tác','CSRF, SQL tham số, escape EJS, ownership/role, rate limit','Integration kiểm tra cả request sai quyền và dữ liệu bất hợp lệ'),('Tài nguyên','mem_limit, log rotation, retention metrics/log','Compose và test rotation; cần quan sát dung lượng thực tế'),('Ngoại lệ','cAdvisor privileged; Promtail root + DAC_READ_SEARCH','Chỉ mục đích quan sát, công khai trong báo cáo')],[3.0,6.4,6.6])
p('Session được lưu PostgreSQL thay vì chỉ giữ trong RAM app. Sau đăng nhập/đăng ký, session được thay mới để giảm session fixation; cookie có thời hạn hai giờ, HttpOnly và SameSite strict. CSRF token được kiểm tra với so sánh thời gian hằng và định dạng cố định. Rate limit đăng nhập là 15 lần trong 15 phút; các request nhạy cảm còn có giới hạn riêng. Cần tránh chạy test đăng nhập liên tiếp rồi nhầm phản hồi 429 là lỗi mật khẩu.')
h('4.2. Quản lý secrets quyền DB và firewall Windows')
p('Mỗi checkout tạo secret ngẫu nhiên riêng: mật khẩu PostgreSQL, role app/monitor, admin, pgAdmin, Grafana và SESSION_SECRET. .env bị gitignore; private key và dump DB cũng vậy. Không đưa mật khẩu vào ảnh commit, ví dụ lệnh shell hoặc báo cáo. Đổi biến môi trường sau khi một volume đã khởi tạo không tự đổi mọi tài khoản; phải thay theo quy trình từng dịch vụ. Lệnh reset-admin chỉ thực hiện có chủ đích vì thu hồi phiên và vô hiệu admin cũ.')
p('Role ứng dụng chỉ có quyền cần thiết để đọc/ghi nghiệp vụ và session, không sở hữu schema. Migration chạy bằng quyền triển khai khi cần tạo/sửa cấu trúc. Role monitoring phục vụ exporter, không dùng chung mật khẩu app. Không mở 5432 lên host giúp giảm một đường truy cập nhưng không thay quyền DB: nếu ứng dụng có lỗi SQL hoặc secret bị lộ, attacker vẫn có thể dùng quyền app trong giới hạn được cấp.')
p('Host đang Windows, vì vậy báo cáo không nhận đã cấu hình UFW/iptables như một host Linux. Get-NetFirewallProfile ghi nhận Domain, Private và Public Enabled; các policy mặc định trả NotConfigured nên không suy ra quy tắc hiệu lực chỉ từ các trường đó. Các port Compose bind 127.0.0.1, chưa thực hiện nghiệm thu từ một máy LAN khác. Không tắt Windows Firewall để xử lý lỗi Docker hoặc mở rộng port cho tiện demo.')
h('4.3. Quét image và khắc phục cảnh báo bảo mật')
p('Sau khi được người dùng cho phép gửi metadata/SBOM tới Docker Scout, image ban đầu được quét với bộ lọc Critical/High. Scout nhận diện 390 package, báo 0 Critical và 13 High thuộc 6 gói: util-linux, perl, zlib và các package npm brace-expansion, undici, http-cache-semantics. Kiểm tra filesystem cục bộ cho thấy các package npm bị báo nằm trong bộ npm của base image; brace-expansion thực sự trong dependency ứng dụng là 2.1.7. Không kết luận rằng npm audit production bằng 0 đồng nghĩa image không có lỗ hổng.')
p('Bản vá chuyển sang multi-stage Alpine 3.23, giữ npm ở build stage và chỉ chép Node cùng dependency production sang runtime, theo hướng dẫn Node Docker [10]. Các gói Debian gây cảnh báo không còn thuộc runtime, npm/Yarn cũng không được đưa sang. Cơ sở dữ liệu package Alpine được giữ nguyên để scanner vẫn nhận diện OS package; không xóa metadata hoặc suppress cảnh báo để làm đẹp kết quả. Một ứng viên Alpine 3.24 còn cảnh báo musl có thông tin fix chưa khớp kho nên không được promote.')
table('So sánh kết quả quét image trước và sau khắc phục',['Chỉ tiêu','Image ban đầu','Image đã triển khai'],[
('Critical / High','0 / 13','0 / 0'),('Medium','Chưa đo do bộ lọc lần đầu','1: BusyBox CVE-2025-60876'),('Low / Unspecified','Chưa đo','0 / 0'),('Package nhận diện','390','141'),('Runtime','Node trên Debian bookworm-slim, có npm','Node 24.22.1 + Alpine 3.23, không npm/Yarn'),('Phạm vi quét','Image app/migrate','Image app/migrate; chưa toàn bộ image hạ tầng')],[4.1,5.9,6.0])
code('docker scout cves local://manh-appointment:local --format markdown `\n  --output docs/evidence/image-scan-remediated.log\n# Lần sau không dùng bộ lọc severity hoặc suppression.')
p('Bản sau được quét mọi severity và còn một Medium của BusyBox 1.37.0-r30, CVE-2025-60876, trạng thái Not Fixed theo dữ liệu Scout cho Alpine 3.23 tại thời điểm kiểm tra. Cảnh báo này chưa được nhận là đã sửa; chưa phân loại đầy đủ khả năng khai thác chỉ từ kết quả scanner. Image mới có manifest sha256:3340102077cf23cc7d0cc18bca1d0afc32fca3d99a050acb814f260cb596e1e3. Kết quả 0 Critical/High không phải chứng nhận an toàn tuyệt đối hoặc áp dụng cho các image hạ tầng chưa quét.')
p('Trước khi thay bản chính, ứng viên được dựng đủ 13 service với project, port và volume riêng; 4 unit test, 3 nhóm integration, smoke/TLS và 16 kiểm tra stack đều đạt. Sau backup, đúng image đã quét được promote, app/migrate được recreate bằng --no-build. Số dòng nghiệp vụ trước/sau không đổi, UID vẫn 1000, runtime không có npm và web/TLS phục hồi. Image cũ được giữ để rollback.')
h('4.4. Dữ liệu bền vững backup và kiểm tra restore')
p('Volume PostgreSQL tách khỏi vòng đời container, nên recreate app hoặc DB container không đồng nghĩa xóa dữ liệu. Tuy nhiên volume trên cùng máy không phải bản sao lưu chống mất máy. scripts/backup.ps1 dùng pg_dump dạng custom, chép dump về backups và tính SHA256; thư mục này không đưa lên Git. File dump tồn tại mới chứng minh đã tạo file, chưa đủ để khẳng định phục hồi được.')
p('Một bản backup appointment-20261024-051045.dump đã được restore vào DB riêng restore_verify_20261024051048. Kết quả giữ schema 1/2/3 và 26 lịch tại thời điểm backup; số dòng và checksum tám bảng nghiệp vụ khớp. Trước khi triển khai vá image đã tạo thêm appointment-20261024-074318.dump với SHA256 99602FD3D7DB2A8735FFD7ED8548F495B706CE05480B5CA8ED47951FFA18D779. Bản mới có checksum file, nhưng không tự nhận đã restore thử bản này nếu chưa có kết quả tương ứng.')
code('powershell -ExecutionPolicy Bypass -File scripts/backup.ps1\npowershell -ExecutionPolicy Bypass -File scripts/restore-test.ps1 `\n  -BackupFile backups/appointment-YYYYMMDD-HHMMSS.dump')
p('Thử restore luôn dùng DB tách biệt, không ghi đè DB chính để lấy ảnh minh chứng. Khi xử lý sự cố phải xác định mất dữ liệu hay chỉ lỗi ứng dụng/mạng. Không dùng docker compose down -v như bước sửa lỗi mặc định vì nó xóa volume. Chính sách sao lưu định kỳ, bản sao ngoài máy, thời gian phục hồi và mức mất dữ liệu chấp nhận được còn cần được thiết kế nếu chuyển sang vận hành thực tế.')
h('4.5. Release rollback và giới hạn tương thích schema')
p('Image ứng dụng có thể được chọn qua APP_IMAGE; mặc định manh-appointment:local. Trước khi thay image, cần lưu tag tốt, image ID, commit source và backup. Tag chỉ là tên có thể đổi, không phải bằng chứng source hoặc backup ngoài máy. Khi muốn chạy đúng image đã chốt phải dùng --no-build; nếu build lại có thể ghi đè tag bằng source hiện tại.')
code('$env:APP_IMAGE=\'manh-appointment:release-0cf3a29\'\ndocker compose up -d --no-deps --no-build --wait app\nnode scripts/verify-stack.js\nnode scripts/verify-tls.js\nRemove-Item Env:APP_IMAGE')
p('Diễn tập rollback đã tạo một lịch fixture đã hủy trong DB thử, ghi số dòng và MD5 nội dung; image ứng viên có CMD thoát lỗi bị readiness từ chối. Chuyển lại image tốt phục hồi HTTP 200, số dòng và checksum lịch khớp trước/sau. Chỉ container/network của project thử bị dừng, volume được giữ; không dùng DB chính để gây lỗi. Đây là bằng chứng rollback code với schema tương thích, không phải thử nâng/hạ mọi schema.')
p('Nếu migration phá tương thích, đổi image không khôi phục cấu trúc hoặc dữ liệu đã thay. Quy trình cần dừng nhận ghi, restore backup vào DB thử, kiểm tra tương thích và đánh giá dữ liệu mới sau backup trước khi quyết định phục hồi DB chính. Phiên bản hiện dùng migration cộng thêm; không thực hiện destructive downgrade trên dữ liệu chính và không nhận đã có phục hồi thảm họa trên máy khác.')

chapter('CHƯƠNG 5. KIỂM THỬ ĐÁNH GIÁ VÀ QUẢN LÝ MÃ NGUỒN')
h('5.1. Phương pháp kiểm thử và phạm vi bằng chứng')
p('Kiểm thử được chia ba tầng. Unit test kiểm tra hàm xử lý dữ liệu độc lập. Integration test gửi HTTP qua Nginx và làm việc với PostgreSQL thật để kiểm tra tương tác nhiều thành phần. Nghiệm thu hạ tầng kiểm tra cấu hình, trạng thái service, quyền, metric, dashboard và LogQL. Mỗi tầng trả lời một câu hỏi khác nhau, vì vậy số test pass không được cộng đơn giản thành điểm đánh giá hoặc bảo đảm không còn bug.')
p('Bộ test tạo dữ liệu giả, dùng email example.com và xử lý lịch thử khi kết thúc. Lịch sử được giữ để kiểm tra trạng thái và snapshot, không xóa dữ liệu chỉ để ảnh trông ít dòng. Test được chạy có pacing để không vô tình vượt rate limit. Các JSON/log kết quả chứa thời điểm và kiểm tra cụ thể được lưu cục bộ; ảnh và cấu hình public được chọn để không lộ thông tin riêng, mật khẩu hoặc token.')
table('Các bộ kiểm tra và kết quả đã đạt',['Bộ kiểm tra','Kết quả','Điều được kiểm chứng'],[
('Unit','4/4','Mật khẩu salted, validation, trạng thái/thời gian, token/ngày Việt Nam'),('Integration PostgreSQL','3/3 nhóm','Đặt lịch đồng thời/idempotency; tài khoản và ownership; đầy đủ quản trị'),('verify-stack.js','16/16','Engine/config/service/job; proxy, UID, DB, exporter, targets, dashboard, LogQL'),('audit-requirements.js','9/9 tại mốc đã push','GitHub, schema/quyền, hardening, headers, từng panel có mẫu và nguồn log'),('verify-tls.js','Đạt TLSv1.3','CA/SAN, HTTP 200, Secure cookie, metrics public 404'),('Clone sạch','Đạt','Clone remote, secret/volume riêng, dựng đủ stack và kiểm tra'),('Backup/restore; rollback','Đạt trong phạm vi thử','Restore DB riêng có checksum; rollback code giữ fixture'),('Image Scout sau vá','0 Critical/High; còn 1 Medium','Mọi severity của image app/migrate, không suy rộng toàn hạ tầng')],[4.4,3.6,8.0])
h('5.2. Các ca kiểm thử nghiệp vụ quan trọng')
table('Ca kiểm thử kỳ vọng và kết quả',['Tình huống','Kỳ vọng và kết quả quan sát'],[
('Hai request cùng slot','Chỉ một lịch giữ slot; request xung đột bị từ chối, DB không có hai lịch active cho slot'),('Hai chuyên viên cùng giờ','Cho phép nếu mỗi người có slot hợp lệ và đúng phân công'),('Gửi lại idempotency key','Cùng nội dung trả lịch cũ; đổi nội dung bị từ chối'),('Khách truy cập lịch người khác','Không được xem/hủy chỉ bằng việc đoán ID'),('Thiếu/sai CSRF hoặc chưa có quyền admin','Thao tác bảo vệ bị từ chối; không đổi danh mục/trạng thái'),('Ngừng đối tượng còn lịch tương lai','Từ chối để bảo vệ lịch đang giữ; lịch sử vẫn tồn tại'),('Slot overlap hoặc batch có lỗi','Từ chối overlap; rollback cả lô nếu một phần không hợp lệ'),('Sửa tên/giá danh mục sau đặt','Lịch giữ snapshot; không viết lại thông tin lịch sử'),('Khóa tài khoản và hồ sơ','Khóa làm mất quyền phiên khách; khóa không tự hủy lịch'),('Xác nhận/hoàn tất sai thời gian','Xác nhận chỉ trước bắt đầu; hoàn tất chỉ sau kết thúc')],[5.6,10.4])
p('Những ca trên quan trọng vì chúng tác động trực tiếp đến tính đúng của lịch hẹn. Frontend có thể hiển thị slot trống nhưng dữ liệu thay đổi ngay trước khi request đến server, vì vậy phải kiểm tra xung đột bằng request đồng thời và DB. Tương tự, một nút bị ẩn không chứng minh quyền được bảo vệ; test phải tự gửi request không đủ quyền. Kết quả hiện tại là các nhóm ca đã triển khai và pass, chưa phải kiểm thử thâm nhập hoặc kiểm thử mọi đường đi của chương trình.')
h('5.3. Các lỗi phát hiện và cách khắc phục có kiểm chứng')
table('Nhật ký lỗi nguyên nhân giải pháp và kiểm tra lại',['Hiện tượng','Nguyên nhân / cách sửa','Kiểm chứng'],[
('Nhầm đăng nhập admin/khách','Tách luồng/username cũ gây nhầm; hợp nhất /login, dùng email và role từ DB','Integration đăng nhập/admin/khách, legacy redirect'),('Proxy lỗi sau recreate app','Nginx giữ IP cũ; dùng Docker DNS resolver với backend là tên dịch vụ','Recreate app, smoke phục hồi không reload Nginx'),('Log rotation sidecar exited 137','Không vào được PID namespace sau restart Nginx; đưa helper cùng container','Xoay log thật và restart, web tiếp tục 200'),('Log mới bị đếm trùng / thời gian sai','Đọc cả file xoay và ingestion timestamp; chỉ đọc active file, parse time JSON','Đối chiếu request_id trước/sau rotation'),('LogQL fail ngay sau request thử','Loki/Promtail chưa ingest; consume body và retry có giới hạn, cập nhật query end','verify-stack không pass cho đến khi có kết quả'),('13 High ở runtime image','Debian và tooling npm không cần tại runtime; build nhiều stage trên Alpine 3.23','Quét mọi severity: 0 High; test candidate và deploy')],[4.2,7.4,4.4])
p('Không xóa log hoặc sửa kết quả failed thành passed bằng tay. Log trùng cũ trong Loki được giữ đến retention và được ghi là giới hạn lịch sử. Với các lỗi vòng đời container, cần kiểm tra sau restart/recreate, vì lệnh cấu hình hợp lệ hoặc lần startup đầu thành công chưa bao phủ hành vi đó. Các commit sửa lỗi được liệt kê ở mục Git để đối chiếu thay đổi thực tế.')
h('5.4. Khả năng tái lập từ GitHub và nghiệm thu deploy')
p('Diễn tập checkout sạch clone remote commit 29cb15c vào thư mục mới, tạo secret ngẫu nhiên và project appointment_fresh_* với volume riêng. Các cổng thử 18080/18443/15050/13001/19090 tránh trùng bản chính. Đủ 13 service được dựng, smoke HTTP, TLS và 16 kiểm tra stack đạt. Sau thử chỉ dừng container/network riêng; không xóa volume hoặc đổi dữ liệu bản chính. Đây là kiểm chứng toàn stack ở mốc cụ thể, khác với chỉ thử migration trên DB rỗng.')
p('Sau thay đổi image bảo mật, ứng viên cuối được kiểm chứng lại bằng test-security-candidate.js trên cùng môi trường thử, nhưng dùng image đã build mới và cấu hình Compose hiện hành. Test chạy đủ unit/integration, smoke, TLS và 16 mục stack; sau đó bản chính được backup và promote đúng image ứng viên. Không nhận đã clone lại toàn bộ báo cáo mới chỉ vì bản vá runtime đã pass; báo cáo ghi rõ mốc của từng loại kiểm chứng.')
p('Địa chỉ dùng demo là HTTPS localhost:8443; công cụ quản trị vẫn có cổng loopback riêng. Trình duyệt cần tin chứng chỉ local theo quyết định của người dùng. Bản 3300 là preview cũ, không dùng để nghiệm thu Nginx, Grafana hay Loki. Một trang HTML mở được ở preview không chứng minh toàn bộ Compose đã chạy; các kết quả deploy được lấy từ bộ nghiệm thu stack và công cụ trên DB thật.')
h('5.5. Lịch sử GitHub và các mốc commit triển khai')
fig(E/'github-commits-latest.png','Lịch sử commit hiển thị trên GitHub của repository đề tài','Ảnh chụp trực tiếp từ trang Commits của repository, thể hiện SHA rút gọn và nội dung các mốc kiểm thử, bảo mật và báo cáo. Bảng bên dưới được trích từ git log và chốt tại một HEAD cụ thể; ảnh chụp có thể sớm hơn commit biên soạn báo cáo chi tiết. Có thể mở từng SHA trên GitHub để xem file thay đổi và nội dung diff.',source='GitHub repository vzmanhdev/dtc245200223_dat_lich_hen, trang Commits được kiểm tra ngày 24/10/2026')
p('Mã nguồn được quản lý tại https://github.com/vzmanhdev/dtc245200223_dat_lich_hen.git. Tên repo chứa MSSV dtc245200223; tài khoản GitHub là vzmanhdev. Bảng bên dưới giữ nguyên tên author và nội dung được ghi trong từng commit, có thể khác tên tài khoản GitHub. Thời gian được chuyển về UTC+7; cách GitHub nhóm ngày có thể tùy múi giờ giao diện.')
p(f'Bảng liệt kê toàn bộ {len(commits)} commit đang có tại HEAD {head[:7]}, trước khi commit bản báo cáo chi tiết này. Nội dung commit giữ nguyên để có thể tìm trực tiếp trên GitHub; không coi tên commit là kết quả kiểm thử nếu thiếu bằng chứng runtime. Các mốc Nginx, monitoring và logging là các thay đổi tách biệt, đáp ứng yêu cầu tối thiểu ba commit có ý nghĩa.')
commit_rows=[]
for i,(full,short,date,author,title) in enumerate(commits,1):
    dt=datetime.fromisoformat(date).astimezone(timezone(timedelta(hours=7)))
    commit_rows.append((i,short,dt.strftime('%d/%m\n%H:%M'),author,title))
table('Danh sách commit thực tế theo thứ tự thời gian',['STT','SHA','UTC+7','Tác giả','Nội dung commit nguyên văn'],commit_rows,[1.2,1.9,1.7,2.5,8.7])
table('Liên hệ các mốc commit với nhóm yêu cầu triển khai',['Nhóm yêu cầu','Commit tiêu biểu','Tác dụng'],[
('Web + PostgreSQL','08fe70d; 17d77bf; fb4e5d6','Nền ứng dụng, tài khoản/chuyên viên, đăng nhập chung và quản trị đầy đủ'),('Nginx / Compose','a2d5e5a; 7d23576; 29cb15c','Gateway, phục hồi DNS sau recreate và HTTPS'),('Monitoring','726abe3; 29cb15c','Datasource/dashboard, exporter và Network I/O'),('Logging','5bec305; 034f0e3; 78b2991; 4280cfa','Pipeline, rotation, timestamp/không trùng và ingestion check'),('Vận hành / hardening','8b9e033; 0cf3a29','Clone sạch/rollback và giảm bề mặt runtime image'),('Báo cáo / minh chứng','9716c89; faefb76; 243d387; 6eb66f3','Đối chiếu, báo cáo ban đầu, biên bản vá và sơ đồ chỉnh lý')],[4.0,5.3,6.7])
p('Tags v1.0.0-local và v1.0.1-local lưu các mốc bàn giao; bản sau có vá image. Code hiện còn có tag image local release-0cf3a29 để quay lại đúng bản đã kiểm thử. Chưa có pipeline CI/CD GitHub Actions trong phạm vi hiện tại, vì vậy ảnh tab Actions hoặc trạng thái push không được dùng để nhận đã triển khai tự động. Thao tác commit/push cuối cùng của báo cáo phát sinh sau mốc chốt bảng và sẽ xuất hiện trên GitHub khi bàn giao.')
h('5.6. Đánh giá mức đáp ứng và các giới hạn còn lại')
p('Các yêu cầu chính của đề 40 đã có sản phẩm và bằng chứng tương ứng: website có đủ dịch vụ, khung giờ, khách hàng và xác nhận; PostgreSQL/pgAdmin; Compose; Nginx với HTTPS và headers; Prometheus/Grafana; Loki/Promtail/LogQL; nhiều biện pháp hardening; source/config/README và lịch sử Git. Báo cáo chi tiết bổ sung mục lục, danh mục hình/bảng, giải thích công cụ, sơ đồ đúng và kết quả bảo mật mới. Kết luận này mô tả hiện trạng trong phạm vi đã kiểm chứng, không phải dự báo điểm chấm.')
p('Còn các giới hạn: nhân viên chưa có tài khoản riêng; chưa xác minh email, quên mật khẩu, nhắc lịch qua email/SMS, thanh toán và đồng bộ lịch ngoài; danh sách quản trị có giới hạn số dòng và bộ lọc, chưa phân trang đầy đủ. Chưa thử tải dài hạn, HA, kiểm thử thâm nhập hoặc toàn bộ image hạ tầng. Một Medium BusyBox chưa có bản vá theo Scout. HTTPS là chứng chỉ tự ký; chưa tự động tin CA trên host. Log tập trung thu ba nguồn cốt lõi, không phải mọi stdout/stderr; lịch sử Loki có thể còn bản trùng cũ.')
p('So với mẫu, báo cáo không nhận đã public Internet ở cổng 80/443, cấu hình UFW trên host Linux hoặc tách mọi nghiệp vụ thành microservices. Thông tin repo/username được trình bày đúng theo tài khoản người dùng cung cấp. Các điểm khác này không được che bằng sơ đồ hoặc ảnh minh họa chung. Việc đánh giá cuối cần kết hợp source, demo và yêu cầu riêng của lớp.')

chapter('KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN')
p('Đề tài đã xây dựng và triển khai một hệ thống đặt lịch hẹn cho đơn vị tư vấn nhiều chuyên viên, có tài khoản khách, phân quyền quản trị, quản lý danh mục và vòng đời lịch. Các ràng buộc chống trùng, transaction, idempotency và snapshot giúp bảo toàn ý nghĩa dữ liệu khi nhiều request hoặc thay đổi danh mục diễn ra. Frontend được kết nối với backend và PostgreSQL thật, không chỉ là giao diện mẫu.')
p('Về triển khai, Docker Compose mô tả toàn bộ môi trường, Nginx phục vụ HTTPS tự ký và reverse proxy, pgAdmin quản trị DB; Prometheus/Grafana và Loki/Promtail hỗ trợ quan sát. Hardening, backup/restore và rollback được kiểm tra ở phạm vi lab. Quá trình quét image phát hiện các rủi ro nằm ngoài dependency ứng dụng và dẫn tới thay đổi runtime; kết quả mới loại bỏ 13 High nhưng vẫn công khai Medium còn lại. Các lỗi DNS proxy, vòng đời rotation và timestamp log được ghi lại cùng kiểm thử tái lập.')
p('Hướng phát triển ưu tiên là tách tài khoản nhân viên chỉ xem lịch được giao, xác minh email và nhắc lịch, chính sách nghỉ/đổi lịch và phân trang. Nếu chuyển sang Internet cần CA tin cậy, giới hạn truy cập các UI quản trị, collector log đang được duy trì, kiểm thử tải/bảo mật, backup tự động ngoài máy và diễn tập phục hồi định kỳ. CI/CD có thể tự động hóa build, scan và test, nhưng cần gate chất lượng và kế hoạch rollback thay vì chỉ tự động chạy lệnh deploy.')

chapter('TÀI LIỆU THAM KHẢO')
refs=[
'[1] Đề tài Thực hành – Triển khai và Quản trị Hệ thống Phần mềm, Đề 40. Tài liệu học phần do người dùng cung cấp; file nguồn trong docx/.',
'[2] Bao_Cao_Trien_Khai_He_Thong.docx. Khung báo cáo triển khai: kiến trúc, web/DB/proxy, metrics/logging và hardening. Không dùng kết luận mẫu thay cho kết quả kiểm tra.',
'[3] Bao_cao_Appointment_System.docx. Tham khảo bố cục trang bìa và cấu trúc báo cáo; không sao chép kết quả/screenshot của người khác. Chapter 01–09 dùng làm tài liệu học phần bổ trợ.',
'[4] Docker Documentation. Networking in Compose. https://docs.docker.com/compose/how-tos/networking/ (truy cập 24/10/2026).',
'[5] PostgreSQL 17 Documentation. Transaction Isolation. https://www.postgresql.org/docs/17/transaction-iso.html (truy cập 24/10/2026).',
'[6] Nginx Documentation. Configuring HTTPS servers. https://nginx.org/en/docs/http/configuring_https_servers.html (truy cập 24/10/2026).',
'[7] MDN Web Docs. X-XSS-Protection header. https://developer.mozilla.org/en-US/docs/Web/HTTP/Reference/Headers/X-XSS-Protection (truy cập 24/10/2026).',
'[8] Prometheus Documentation. Overview. https://prometheus.io/docs/introduction/overview/ (truy cập 24/10/2026).',
'[9] Grafana Loki Documentation. Promtail agent. https://grafana.com/docs/loki/latest/send-data/promtail/ (truy cập 24/10/2026).',
'[10] Node.js Docker Working Group. Docker and Node.js Best Practices, Smaller images without npm/yarn. https://github.com/nodejs/docker-node/blob/main/docs/BestPractices.md (truy cập 24/10/2026).',
'[11] Repository dự án và lịch sử commit. https://github.com/vzmanhdev/dtc245200223_dat_lich_hen. Source/config cùng kết quả kiểm thử cục bộ ngày 24/10/2026 là nguồn đối chiếu hiện trạng triển khai.'
]
for ref in refs:p(ref)

chapter('PHỤ LỤC')
h('A. Bộ lệnh dựng và nghiệm thu')
code('docker compose config --quiet\ndocker compose up -d --build --wait --wait-timeout 240\ndocker compose ps -a\ndocker compose exec -T nginx nginx -t\ndocker compose exec -T app id -u\nnpm.cmd test\nnpm.cmd run test:integration\nnode scripts/verify-stack.js\nnode scripts/verify-tls.js\nnode scripts/audit-requirements.js\ngit log --oneline --decorate -25\ngit status --short')
p('Chạy từ thư mục repo có .env cục bộ và Docker Engine đang hoạt động. verify-stack cần đủ thời gian scrape/ingestion; integration tạo dữ liệu giả và không nên chạy lặp sát nhau vì giới hạn đăng nhập. audit-requirements đối chiếu HEAD với GitHub, vì vậy commit chưa push có thể làm mục đó fail; không thay kết quả bằng cách bỏ điều kiện.')
h('B. Danh mục route và API chính')
table('Route phương thức quyền và tác dụng',['Phương thức / route','Quyền','Tác dụng'],[
('GET /; /login; /register','Công khai','Xem giao diện, đăng nhập thống nhất, đăng ký khách'),('POST /login; /register; /logout','Theo session/CSRF','Xác thực/tạo khách/hủy phiên; có rate limit ở login/register'),('GET /api/services; /api/staff; /api/slots','Công khai','Đọc danh mục và lịch trống theo điều kiện'),('POST /api/appointments','Khách + CSRF','Đặt lịch trong transaction, kiểm tra idempotency/slot'),('GET /account; /api/my/appointments','Khách','Giao diện/dữ liệu lịch thuộc tài khoản'),('POST /api/my/appointments/:id/cancel','Khách + CSRF','Hủy lịch sở hữu trước giờ bắt đầu'),('POST /api/my/profile','Khách + CSRF','Đổi tên/điện thoại hồ sơ'),('GET /admin[/services|/staff|/slots|/customers]','Admin','Năm phân hệ quản trị; có bộ lọc/chỉnh sửa'),('POST /api/admin/appointments/:id/status','Admin + CSRF','Xác nhận/hủy/hoàn tất với quy tắc thời gian'),('GET /metrics; /health/*','Nội bộ','Metrics/readiness; public qua Nginx trả 404')],[7.3,3.2,5.5])
p('Các mutation danh mục được định nghĩa trong management.js và nhận POST kèm CSRF, chạy transaction. Bảng trên là danh mục rút gọn phục vụ trình bày, không tự nhận mọi route là REST đầy đủ. Legacy /admin/login chuyển về /login; token lookup/cancel còn phục vụ lịch cũ, không thay quyền tài khoản cho luồng đặt mới.')
h('C. Kịch bản demo khoảng mười lăm phút')
table('Kịch bản thao tác và điều cần chỉ ra',['Thời gian','Thao tác','Kết quả phải đối chiếu'],[
('0–1 phút','Giới thiệu bài toán; ps -a','11 service chạy và hai job exit 0'),('1–4','Admin thêm dịch vụ/nhân viên, phân công, mở slot','Thời lượng đúng, không overlap cùng người'),('4–6','Khách đăng nhập và gửi yêu cầu','Mã lịch mới, pending, đúng người/giờ'),('6–8','Admin xác nhận, khách tải lại account','Cùng mã chuyển confirmed; quyền khách được giữ'),('8–9','pgAdmin xem bảng liên quan','Dữ liệu và sự kiện trong PostgreSQL thật'),('9–11','Targets và Grafana','5 job UP, web/DB/CPU/RAM/I/O có dữ liệu'),('11–13','Explore ba LogQL','Log thật, parse JSON, số đếm và phạm vi thời gian'),('13–15','HTTPS, hardening, Git và backup','TLS/Secure, UID/quyền, lịch sử commit và giới hạn')],[2.7,6.1,7.2])
p('Mở HTTPS trước khi trình bày và tự xử lý việc tin chứng chỉ local nếu người demo đồng ý. Không hiển thị .env hoặc mật khẩu lên màn chiếu. Dùng khách demo và ngày tương lai, không chỉnh dữ liệu người dùng thật để minh họa. Kịch bản không thay video nộp nếu lớp yêu cầu quay; tại thời điểm viết chưa có video ghi hình trong hồ sơ.')
h('D. Chú giải nguồn minh chứng và cách kiểm tra lại')
p('Ảnh pgAdmin, Grafana, Explore và giao diện được lấy từ docs/evidence; chú thích dưới từng hình giải thích vùng cần đọc và giới hạn thời điểm chụp. Các sơ đồ mới được lưu cả PNG độ phân giải cao và SVG có thể chỉnh sửa, được dựng theo compose.yaml, schema và source. Đường mũi tên được quy ước riêng cho từng sơ đồ; không dùng một loại mũi tên vừa biểu diễn request vừa biểu diễn dữ liệu mà không có chú giải.')
note('Các JSON vận hành lưu cục bộ; ảnh, config và báo cáo lưu trong Git. Tài liệu gốc của học phần được giữ nguyên.')

# Remove unused images and hyperlinks from the other student's body.
used=set()
for el in d.element.iter():
    for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
        if el.get(attr):used.add(el.get(attr))
for rid,rel in list(d.part.rels.items()):
    if rid not in used and rel.reltype.rsplit('/',1)[-1] in ['image','hyperlink']:d.part.drop_rel(rid)
settings=d.settings.element;u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
OUT.parent.mkdir(exist_ok=True);d.save(OUT)
(QA/'manifest.json').write_text(json.dumps({'output':str(OUT),'reference_sha256':sha256(REF.read_bytes()).hexdigest(),'git_head':head,'commit_count':len(commits),'figures':figures,'tables':tables,'paragraphs':len(d.paragraphs)},ensure_ascii=False,indent=2),encoding='utf8')
print('Created',OUT,'figures',len(figures),'tables',len(tables),'commits',len(commits))
