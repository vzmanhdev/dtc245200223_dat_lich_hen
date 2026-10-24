"""Build the student's report from the retained cover; run with python-docx + Pillow.
Original course documents are read-only and intentionally not distributed in Git.
"""
from pathlib import Path
from hashlib import sha256
from copy import deepcopy
from zipfile import ZipFile
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
REF = ROOT/'docx/Bao_cao_Appointment_System.docx'
OUT = ROOT/'reports/Bao_cao_Appointment_dtc245200223.docx'
TMP = ROOT/'tmp/report-qa'
TMP.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(exist_ok=True)
fingerprint = sha256(REF.read_bytes()).hexdigest()
(TMP/'artifact.md').write_text(f'''# Report execution contract
Reference: {REF}; SHA256 {fingerprint}; 35 pages rendered with Microsoft Word.
User requested page 1 cover only. Original reference stays unchanged.
Cover: retain document.xml body paragraphs 0..14, including logo drawing,
double page border, center university/faculty/title blocks, right date and
section break in paragraph 14. A4 21.006 x 29.693 cm; top/bottom 2 cm,
left 3 cm, right 2 cm; Times New Roman. Change student/name/class/date slots
in paragraphs 7,8,14 only. Teacher retained from source, not independently confirmed.
Body is an authorized replacement of the sample's results, all old body removed.
Reuse retained styles Heading 1/2, Body Text, footer PAGE and second section
A4 21.001 x 29.700 cm, margins top2.3495/bottom2.1996/left3/right2 cm.
Body 13pt Times New Roman, 1.15 lines, 6pt after; headings 16/14 bold black;
tables 11.5pt, header repeats, no row splitting; figures fit text width16cm.
Preserve cover geometry, logo, border, footer parts; remove unused sample image
relationships to avoid carrying other students' screenshots in the output.
Update fields using installed Microsoft Word, then inspect every final page.
Reference package inventory: tmp/report-reference/inventory.txt.
Source cover visual: tmp/report-reference/cover.png.
''', encoding='utf8')

def diagram(name, boxes, arrows, size=(1400, 740)):
    im=Image.new('RGB', size, 'white'); dr=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
    small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
    for a,b in arrows:
        dr.line([a,b],fill='#274d43',width=4)
        import math
        angle=math.atan2(b[1]-a[1],b[0]-a[0])
        dr.polygon([b,(b[0]-17*math.cos(angle-.45),b[1]-17*math.sin(angle-.45)),(b[0]-17*math.cos(angle+.45),b[1]-17*math.sin(angle+.45))],fill='#274d43')
    for rect,title,lines in boxes:
        dr.rounded_rectangle(rect,radius=12,fill='#edf3e9',outline='#274d43',width=3)
        x0,y0,x1,y1=rect; dr.text((x0+18,y0+14),title,font=font,fill='#173e34')
        for i,line in enumerate(lines):dr.text((x0+18,y0+57+i*31),line,font=small,fill='#222222')
    path=TMP/(name+'.png');im.save(path);return path

arch=diagram('architecture',[
((25,40,325,205),'Windows / localhost',['HTTP 8080 / TLS 8443','pgAdmin 5050','Grafana 3001 / Prom 9090']),
((440,40,770,180),'Nginx reverse proxy',['TLS + security headers','internal health / metrics']),
((925,40,1370,180),'Express + EJS',['customer / admin modules','session, CSRF, transaction']),
((925,300,1370,440),'PostgreSQL 17',['business + sessions','internal DB network']),
((440,300,770,470),'Prometheus + Grafana',['app / Nginx / DB exporters','cAdvisor CPU, RAM, I/O']),
((25,510,770,700),'Promtail -> Loki -> Grafana Explore',['read-only app / Nginx / PostgreSQL logs','JSON event time for app and Nginx; bounded retention']),
], [((325,105),(440,105)),((770,105),(925,105)),((1140,180),(1140,300)),((1010,180),(770,320)),((1100,440),(770,560)),((605,180),(400,510))])
erd=diagram('erd',[
((20,30,430,190),'customer_accounts',['id PK, email UNIQUE','password_hash, active']),
((20,300,430,460),'customers',['id PK; contact snapshot','legacy contacts retained']),
((500,230,920,470),'appointments',['id PK, code UNIQUE','account_id / customer_id FK','slot_id FK, status, snapshots','appointments -> events (1:N)']),
((990,230,1380,470),'time_slots',['id PK; staff_id / service_id FK','start_at, end_at, is_open','overlap / held slot constraints']),
((500,30,920,165),'services',['id PK, duration, price, active']),
((990,30,1380,165),'staff + staff_services',['staff (1:N) time_slots','staff (N:M) services']),
((500,565,1380,710),'Integrity rules',['One held appointment per slot; one open time range per staff','Transactions + row locks; preserve service/staff/price snapshots']),
], [((430,110),(500,280)),((430,380),(500,380)),((920,350),(990,350)),((750,165),(1080,230)),((1190,165),(1190,230))])

d=Document(REF)
d.sections[1].different_first_page_header_footer=False
# Cover ends in its original section break; discard all other student's results.
keep={id(p._p) for p in d.paragraphs[:15]}
cover_elements=list(d.element.body)[:15]
for element in list(d.element.body):
    if element not in cover_elements and element.tag!=qn('w:sectPr'): d.element.body.remove(element)
def replace_keep(p,text):
    runs=p.runs
    if runs:
        runs[0].text=text
        for r in runs[1:]:r.text=''
    else:p.add_run(text)
replace_keep(d.paragraphs[7],'      SINH VIÊN THỰC HIỆN:')
replace_keep(d.paragraphs[8],'      HỌ VÀ TÊN: NGUYỄN VĂN MẠNH     LỚP: CNTTK23B\n      MÃ SINH VIÊN: dtc245200223\n\nGIẢNG VIÊN HƯỚNG DẪN: NGUYỄN ANH CHUYÊN')
for r in d.paragraphs[8].runs:r.font.size=Pt(13)
replace_keep(d.paragraphs[14],'Thái Nguyên, ngày 24 tháng 10 năm 2026')
if 'Caption' not in d.styles:
    d.styles.add_style('Caption',WD_STYLE_TYPE.PARAGRAPH).base_style=d.styles['Table Caption']
# Keep the source's named body styles, with explicit black typography.
for name in ['Body Text','First Paragraph','Heading 1','Heading 2','Front Heading','Caption','Table Caption']:
    if name in d.styles:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.line_spacing=1.15
        if name in ['Body Text','First Paragraph']:st.font.size=Pt(13);st.paragraph_format.first_line_indent=Cm(.5)
        if name.startswith('Heading') or name=='Front Heading':
            st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.page_break_before=False
            st.font.size=Pt(16 if name in ['Heading 1','Front Heading'] else 14)
for st in d.styles:
    if st.type==1 and st.name.startswith('Heading'):st.font.color.rgb=RGBColor(0,0,0)
d.core_properties.author='Nguyễn Văn Mạnh'
d.core_properties.last_modified_by='Nguyễn Văn Mạnh'
d.core_properties.title='Triển khai hệ thống đặt lịch hẹn – dtc245200223'
d.core_properties.subject='Đề 40 – Docker Compose localhost'
d.core_properties.comments='Kết quả thực hành trên localhost, 24/10/2026.'
pages=[]
def page(title):
    p=d.add_paragraph(title,'Heading 1');p.paragraph_format.page_break_before=bool(pages)
    pages.append(title)
def h(t):d.add_paragraph(t,'Heading 2')
def p(t):return d.add_paragraph(t,'Body Text')
def code(t):
    q=d.add_paragraph();q.paragraph_format.space_after=Pt(7);q.paragraph_format.line_spacing=1.05
    r=q.add_run(t);r.font.name='Consolas';r.font.size=Pt(9)
def table(headers, rows, widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'999999');borders.append(b)
    t._tbl.tblPr.append(borders)
    if widths:
        for col,w in zip(t.columns,widths):col.width=Cm(w)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();x=OxmlElement('w:cantSplit');trPr.append(x)
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[j])
            for q in cell.paragraphs:
                q.paragraph_format.space_after=Pt(4);q.paragraph_format.space_before=Pt(3);q.paragraph_format.line_spacing=1.05
                for r in q.runs:r.font.name='Times New Roman';r.font.size=Pt(11.5);r.font.bold=(i==0)
    return t
def fig(path, caption, width=16):
    q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Cm(width))
    q=d.add_paragraph(caption,'Caption');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in q.runs:r.font.size=Pt(11);r.font.italic=True
E=ROOT/'docs/evidence'

page('LỜI MỞ ĐẦU VÀ CẤU TRÚC BÁO CÁO')
p('Báo cáo trình bày việc xây dựng và triển khai đề tài số 40: website đặt lịch hẹn cho một đơn vị tư vấn có nhiều chuyên viên. Trọng tâm là đưa ứng dụng có frontend, backend và PostgreSQL vào môi trường Docker Compose, đồng thời thiết lập reverse proxy, cơ sở dữ liệu quản trị, monitoring, logging và hardening theo yêu cầu học phần.')
p('Phạm vi được thống nhất là vận hành trên máy cá nhân Windows với WSL 2 và Docker Desktop. Các cổng chỉ bind 127.0.0.1; không có VPS hoặc domain công khai. Kết quả, ảnh và kiểm thử trong báo cáo thuộc hệ thống của Nguyễn Văn Mạnh, không sử dụng kết quả của báo cáo mẫu. Bìa giữ bố cục trang đầu mẫu theo yêu cầu; tên giảng viên giữ theo mẫu và cần đối chiếu thông tin lớp trước khi nộp.')
table(['Nội dung','Phạm vi trình bày'],[
('Chương 1','Bài toán, tác nhân, chức năng và tiêu chí nghiệm thu'),
('Chương 2','Kiến trúc ứng dụng, dữ liệu và hạ tầng Docker'),
('Chương 3','Sáu bước: source/Compose, Nginx, DB, metrics, logs, hardening'),
('Chương 4','Giao diện, kiểm thử, dựng từ clone sạch, backup/restore'),
('Chương 5','Đánh giá, giới hạn, hướng phát triển và tài liệu tham khảo'),
('Phụ lục','Kịch bản demo và bộ lệnh bàn giao')],[3.4,12.6])
h('Quy ước đọc kết quả')
p('Ảnh là trạng thái tại thời điểm chụp ngày 24/10/2026, không phải số liệu cố định. “Đạt” chỉ mô tả phạm vi đã kiểm chứng; không dự báo điểm chấm. Các mật khẩu nằm trong .env của máy thực hành và không xuất hiện trong báo cáo hoặc Git. Bộ tài liệu chapter của học phần dùng để tham khảo kiến thức; đề 40 là nguồn yêu cầu chính.')

page('CHƯƠNG 1. PHÂN TÍCH BÀI TOÁN')
h('1.1. Bối cảnh và mục tiêu')
p('Khách hàng cần đặt một cuộc tư vấn với đúng chuyên viên, đúng dịch vụ và thời điểm còn trống. Nếu chỉ nhập tên và giờ vào biểu mẫu, đơn vị khó xác định chủ sở hữu lịch, dễ nhận trùng và khó thông báo trạng thái. Hệ thống giải quyết bằng tài khoản khách hàng, danh mục dịch vụ, phân công nhân viên, khung giờ được mở trước và quá trình xác nhận do quản trị viên điều phối.')
p('Khách đặt hẹn với chuyên viên thuộc đơn vị Mạnh Consulting. Một chuyên viên có thể phụ trách nhiều dịch vụ; một dịch vụ có thể có nhiều chuyên viên. Hai chuyên viên có thể nhận khách cùng thời điểm, nhưng một người không được có hai khung giờ đang mở chồng lấn. Giá và thời lượng được quản trị ở dịch vụ, không lấy tùy ý từ trình duyệt.')
table(['Nhóm yêu cầu đề 40','Sản phẩm thực hiện'],[
('Dịch vụ / khung giờ / khách hàng / xác nhận','Ứng dụng Express + EJS, tài khoản, quản trị 5 phân hệ'),
('PostgreSQL + pgAdmin','Schema 3 phiên bản; pgAdmin kết nối DB thật'),
('Docker Compose và Nginx','13 service; reverse proxy; HTTP và TLS tự ký'),
('Prometheus + Grafana','5 target; 15 panel gồm CPU, RAM, Network I/O'),
('Loki + Promtail + LogQL','3 nguồn log cốt lõi và tối thiểu 3 truy vấn'),
('Hardening; GitHub; báo cáo; demo','Nhiều lớp bảo vệ; commit theo mốc; báo cáo và kịch bản')],[6,10])
h('1.2. Tiêu chí hoàn thành')
p('Một lượt demo phải tạo được lịch bằng tài khoản khách, xác nhận ở admin và xem lại trạng thái ở tài khoản đó; PostgreSQL lưu đúng dữ liệu. Đồng thời người xem mở được pgAdmin, dashboard có mẫu dữ liệu, LogQL trả log thật và kiểm tra được cấu hình hardening. Khởi động lại app phải giữ dữ liệu nhờ volume và session PostgreSQL.')

page('1.3. TÁC NHÂN VÀ QUY TẮC NGHIỆP VỤ')
table(['Tác nhân','Quyền và trách nhiệm'],[
('Khách chưa đăng nhập','Xem dịch vụ và lịch trống, đăng ký/đăng nhập; đặt lịch cần tài khoản'),
('Khách có tài khoản','Đặt lịch, xem lịch của mình, hủy trước giờ bắt đầu, cập nhật tên/điện thoại'),
('Quản trị viên','Dịch vụ, nhân viên, phân công, khung giờ, khách hàng, xác nhận/hủy/hoàn tất'),
('Chuyên viên tư vấn','Đối tượng được phân công nhận lịch; admin điều phối, chưa có tài khoản riêng')],[4.2,11.8])
h('Trạng thái lịch hẹn')
table(['Trạng thái','Chuyển trạng thái hợp lệ'],[
('pending – chờ xác nhận','Admin xác nhận trước giờ bắt đầu; khách/admin có thể hủy theo quyền'),
('confirmed – đã xác nhận','Admin hoàn tất sau giờ kết thúc; có thể hủy theo quyền'),
('completed – hoàn thành','Kết thúc, không mở lại bằng thao tác thông thường'),
('cancelled – đã hủy','Kết thúc; slot tương lai còn mở có thể nhận lượt đặt mới')],[5,11])
p('Luồng chuẩn: khách đăng nhập → chọn dịch vụ → chuyên viên được phân công → ngày và giờ còn trống → gửi yêu cầu → pending → admin xác nhận → confirmed → admin hoàn tất sau khi tư vấn kết thúc. Khách hủy chỉ trước giờ bắt đầu; admin có thể hủy lịch quá giờ để xử lý tồn đọng.')
p('Backend kiểm tra vai trò từ DB và ràng buộc quyền sở hữu, không tin ID hay thông tin liên hệ do client tự gửi. Khi dịch vụ/nhân viên còn lịch tương lai đang giữ, hệ thống ngăn thao tác ngừng hoặc sửa phân công gây mất lịch. Đổi thời lượng phải xử lý các slot mở trước, tránh tạo khoảng giờ sai với dịch vụ.')
h('Phi chức năng')
p('Ưu tiên tính đúng dữ liệu, khả năng tái lập, quan sát và phục hồi trong lab. SQL tham số hóa, mật khẩu băm có salt, CSRF, rate limit và transaction bảo vệ luồng nhận lịch. Không đặt mục tiêu SLA, tải lớn hoặc bảo đảm an toàn tuyệt đối khi chưa có kiểm thử tương ứng.')

page('CHƯƠNG 2. KIẾN TRÚC VÀ CÔNG NGHỆ')
p('Ứng dụng là backend Express được chia module, render giao diện EJS và JavaScript phía trình duyệt. Đây là ứng dụng nguyên khối có module, kết hợp nhiều dịch vụ hạ tầng trong Docker; không gọi các module nghiệp vụ là microservices độc lập. HTTP/TLS đi vào Nginx rồi chuyển đến app:3000; app kết nối PostgreSQL qua network nội bộ.')
fig(arch,'Hình 2.1. Kiến trúc triển khai thực tế trên localhost.')
table(['Thành phần','Công nghệ và chức năng'],[
('Frontend / backend','EJS, CSS, JavaScript / Node.js 24, Express 5'),
('Dữ liệu','PostgreSQL 17; pgAdmin; migrations 1/2/3'),
('Reverse proxy','Nginx unprivileged; TLS tự ký; headers; rate limit'),
('Quan sát','Prometheus, exporters, cAdvisor, Grafana, Loki, Promtail')],[4.2,11.8])
p('Có 13 service Compose: 11 service chạy liên tục và hai job log-init/migrate kết thúc exit 0. Healthcheck và depends_on chờ DB, migration và app sẵn sàng trước khi mở proxy. Image dùng digest Linux/amd64; volume giữ dữ liệu khi thay container.')

page('2.2. THIẾT KẾ DỮ LIỆU VÀ TÍNH NHẤT QUÁN')
fig(erd,'Hình 2.2. Sơ đồ quan hệ rút gọn; session và migrations là bảng hỗ trợ.')
p('customer_accounts lưu tài khoản và trạng thái hoạt động; customers giữ liên hệ gắn lịch, bao gồm dữ liệu cũ. services chứa tên, thời lượng và giá; staff chứa hồ sơ chuyên viên; staff_services biểu diễn phân công nhiều–nhiều. time_slots xác định người, dịch vụ và khoảng giờ. appointments liên kết chủ tài khoản, liên hệ và slot; appointment_events lưu lịch sử chuyển trạng thái.')
p('Khi đặt, transaction khóa slot và kiểm tra lại tính hoạt động của dịch vụ/nhân viên. Unique index chỉ cho một lịch pending/confirmed giữ slot; constraint ngăn các khoảng đang mở chồng nhau theo từng nhân viên. Điều kiện này vẫn bảo vệ được khi hai request cùng vượt qua bước kiểm tra ban đầu. Idempotency gắn với tài khoản giúp gửi lại cùng yêu cầu không nhân đôi lịch.')
p('Tên dịch vụ, tên nhân viên, giá và thời lượng có snapshot trong lịch. Sửa danh mục sau này không làm lịch sử đổi nghĩa. Đối tượng được ngừng hoạt động thay vì xóa cứng khi còn quan hệ. Lịch cũ chưa có tài khoản không tự gán theo email chưa xác minh. Migrate v3 giữ dữ liệu và chỉ seed lần đầu; khởi động lại không tái tạo nhân viên đổi tên hoặc slot đã đóng.')
h('Phân quyền cơ sở dữ liệu')
p('Role ứng dụng không phải superuser và không sở hữu schema. Migration dùng quyền riêng để tạo cấu trúc; exporter dùng role giám sát tách biệt. Cổng 5432 không publish ra host. pgAdmin dùng hostname postgres trong Docker network, không dùng localhost của container để kết nối DB.')

page('CHƯƠNG 3. QUY TRÌNH TRIỂN KHAI')
h('Bước 1 – Chuẩn bị source và Docker Compose')
p('Môi trường thực hành dùng Windows, WSL 2 và Docker Desktop chạy Linux containers. Node.js 24 phục vụ script cấu hình/kiểm thử; Git for Windows có OpenSSL để tạo chứng chỉ. Source được quản lý tại repository dưới đây; .env, private key, backup và tài liệu học gốc không đưa lên Git.')
code('git clone https://github.com/vzmanhdev/dtc245200223_dat_lich_hen.git\ncd dtc245200223_dat_lich_hen\nnpm.cmd ci --ignore-scripts\nnpm.cmd run setup\npowershell -ExecutionPolicy Bypass -File scripts/start.ps1')
p('setup tạo mật khẩu ngẫu nhiên riêng trong .env và chứng chỉ tự ký của máy. Nếu .env đã tồn tại, script giữ nguyên; không ghi đè thông tin hiện có. Sau khi volume DB đã được tạo, chỉ sửa biến mật khẩu trong .env không tự đổi mọi tài khoản. Quy trình reset admin là thao tác chủ động riêng, không chạy mỗi lần khởi động.')
table(['Network','Thành phần và mục đích'],[
('frontend','Nginx → app, nhận HTTP/TLS qua loopback'),
('db (internal)','PostgreSQL, app/migrate, pgAdmin, DB exporter'),
('monitoring (internal)','Prometheus, exporters, cAdvisor, Loki, Promtail, Grafana'),
('management','Giao diện quản trị công cụ qua cổng 127.0.0.1')],[4.4,11.6])
h('Kiểm tra trạng thái')
code('docker compose config --quiet\ndocker compose up -d --build --wait --wait-timeout 240\ndocker compose ps -a\nnode scripts/verify-stack.js')
p('Migrate và log-init exited 0 là trạng thái hoàn thành bình thường. Với 11 service chạy dài hạn, exited hoặc restarting không được coi là đạt. start.ps1 tự tìm Docker CLI cài theo người dùng trên Windows và chạy smoke sau khi Compose sẵn sàng.')

page('BƯỚC 2 – NGINX, HTTPS VÀ HEADERS')
p('Nginx tiếp nhận HTTP tại 127.0.0.1:8080 và HTTPS tự ký tại 127.0.0.1:8443. Cả hai chuyển đến cùng backend nên không tạo hai hệ thống đăng nhập. Cấu hình dùng Docker DNS resolver và tên app thay vì giữ một IP cố định, giúp proxy phục hồi khi app được recreate.')
code('listen 8443 ssl;\nssl_certificate /etc/nginx/certs/localhost.crt;\nssl_certificate_key /etc/nginx/certs/localhost.key;\nssl_protocols TLSv1.2 TLSv1.3;\nresolver 127.0.0.11 valid=5s ipv6=off;\nset $appointment_backend app:3000;\nproxy_pass http://$appointment_backend;')
p('Chứng chỉ RSA 2048 có SAN localhost và 127.0.0.1, hiệu lực 365 ngày kể từ lúc tạo. Key/cert được tạo riêng, mount chỉ đọc và gitignore. Cookie session đặt Secure tự động khi request qua HTTPS; HTTP giữ chế độ phù hợp để demo cục bộ. Không bật HSTS cho localhost để tránh ảnh hưởng ứng dụng khác trên cùng hostname.')
table(['Kiểm tra','Kết quả'],[
('nginx -t','Cấu hình hợp lệ'),
('verify-tls.js','TLS 1.3; CA/SAN được kiểm tra; cookie Secure'),
('Header phản hồi','nosniff, SAMEORIGIN, Referrer/Permissions, CSP'),
('Endpoint nội bộ','/metrics, /health/*, /stub_status trả 404 qua cổng public')],[5,11])
p('Kiểm tra TLS kết nối trực tiếp Nginx từ mạng Docker, tin cậy đúng file CA tự ký và giữ rejectUnauthorized=true. Cách này tránh phần mềm antivirus trên Windows thay chứng chỉ localhost. Trình duyệt chưa tin CA có thể cảnh báo; người dùng tự quyết định tin chứng chỉ local. Không tự thêm CA vào Windows và không tắt xác minh TLS.')
p('X-XSS-Protection để 0 theo cơ chế hiện tại; biện pháp chính là CSP và escape nội dung EJS. Không sao chép header đã lỗi thời từ mẫu chỉ để giống văn bản. Nginx cũng giới hạn body 16 KiB, request rate và timeout; các lớp này giảm lạm dụng nhưng chưa thay kiểm thử tải.')

page('BƯỚC 3 – POSTGRESQL VÀ PGADMIN')
p('PostgreSQL lưu nghiệp vụ và session vào volume postgres_data. Migration chạy sau pg_isready; app chỉ khởi động khi migration thành công. Mật khẩu DB tách với app và monitoring, dùng SCRAM-SHA-256. Không mở cổng DB ra Windows; truy cập bằng pgAdmin hoặc docker compose exec khi quản trị.')
fig(E/'pgadmin-services-verified.png','Hình 3.1. pgAdmin truy vấn bảng services; Query complete trên DB thật.')
p('Mở http://localhost:5050 và dùng PGADMIN_DEFAULT_EMAIL/PASSWORD trong .env. Server khai báo sẵn dùng host postgres, database appointment, port 5432; nhập mật khẩu DB cục bộ khi được hỏi. Ảnh chỉ minh họa dữ liệu dịch vụ, không đưa bảng chứa password_hash hoặc thông tin tài khoản riêng vào báo cáo.')
code('SELECT id, name, duration_minutes, price, active\nFROM services ORDER BY id;\nSELECT status, count(*) FROM appointments GROUP BY status;')
p('Số dòng phụ thuộc dữ liệu kiểm thử ở thời điểm chạy. Kiểm tra sâu xác nhận schema 1/2/3, constraint nghiệp vụ và role app không sở hữu schema. Đổi container app không làm mất lịch; dữ liệu chỉ mất nếu volume bị xóa, vì vậy không dùng down -v trong thao tác vận hành thông thường.')

page('BƯỚC 4 – PROMETHEUS VÀ GRAFANA')
p('Prometheus thu 5 job: prometheus, appointment, nginx, postgres và cadvisor. App cung cấp số request, mã trạng thái và histogram thời gian xử lý; Nginx exporter đọc stub_status ở cổng nội bộ; DB exporter đọc số kết nối/giao dịch; cAdvisor quan sát tài nguyên container. Tại lần nghiệm thu, cả 5 job UP và các truy vấn dashboard có mẫu hữu hạn.')
fig(E/'grafana-network-verified.png','Hình 3.2. Network receive/transmit theo Compose service và log app thật.')
p('Dashboard appointment-local được provision từ Git, gồm 15 panel: 14 panel Prometheus và một panel Loki. Có target UP, HTTP rate/status/p95, Nginx, kết nối/giao dịch DB, CPU, RAM và Network I/O. Hai panel mới dùng rate(container_network_receive_bytes_total[2m]) và transmit tương ứng, cộng theo service, bỏ interface lo và hiển thị B/s.')
p('Cần có tối thiểu hai lần scrape để rate có dữ liệu. Giá trị 0 có thể hợp lệ nếu hệ thống ít tải; “No data” không được tự hiểu là đạt. Thử clone độc lập trên cùng Docker host có thể làm cAdvisor thấy thêm container; cần dừng bản thử hoặc lọc project trước khi diễn giải tải của bản chính.')
code('sum(up)\npg_up\nnginx_up\nnode scripts/audit-requirements.js')

page('BƯỚC 5 – LOKI, PROMTAIL VÀ LOGQL')
p('Promtail đọc ba nguồn log cốt lõi: Pino JSON của app, access JSON Nginx và file log PostgreSQL. Mount chỉ đọc, không cấp Docker socket hay volume dữ liệu DB. App không ghi body request, mật khẩu hoặc query string vào log HTTP. Phạm vi này đáp ứng hệ log tập trung của đề; chưa thu toàn bộ stdout/stderr của mọi container.')
fig(E/'loki-logql-verified.png','Hình 3.3. Grafana Explore trả log lỗi HTTP bằng LogQL thực tế.')
code('{service="app", environment="local"}\n\n{service="nginx", environment="local"} | json | __error__="" | status >= 400\n\nsum(count_over_time({service="nginx", environment="local"}\n  | json | __error__="" | status >= 400 [5m]))')
p('Query thứ nhất xem request app; query thứ hai lọc HTTP lỗi, gồm 404 cố ý dùng kiểm thử; query thứ ba đếm sự kiện trong cửa sổ 5 phút. Bộ verify-stack tạo request thiếu trang và kiểm tra ba query có kết quả sau khi Promtail chuyển log. Các số đếm không được coi là SLA hay lỗi nghiệp vụ nếu chưa phân biệt request kiểm thử.')
p('Pipeline app/Nginx parse trường time làm timestamp và chỉ đọc file hiện hành, tránh thu lại log đã xoay. PostgreSQL plaintext dùng thời điểm thu. Bản trùng lịch sử trước khi sửa vẫn có thể còn trong Loki đến hết retention. Promtail giữ theo yêu cầu lab; cấu hình không được quảng bá như giải pháp Internet dài hạn.')

page('BƯỚC 6 – HARDENING VÀ PHẠM VI BẢO VỆ')
table(['Biện pháp','Cách áp dụng và bằng chứng'],[
('Non-root và ít quyền','App UID node; Nginx unprivileged; app cap_drop ALL, no-new-privileges'),
('Filesystem','App/Nginx read_only; tmpfs /tmp giới hạn; chỉ volume cần thiết được ghi'),
('Tách mạng và cổng','db/monitoring internal; không publish 5432; mọi port host 127.0.0.1'),
('Tách quyền DB','App không superuser/owner; monitoring riêng; migration riêng'),
('Secrets','Mật khẩu sinh riêng trong .env; key/cert local; không COPY vào image/Git'),
('Web và xác thực','TLS, headers, cookie HttpOnly/SameSite, Secure qua TLS; session server-side'),
('Bảo vệ thao tác','CSRF, tham số SQL, escape EJS, rate limit, xác minh ownership/role'),
('Giới hạn tài nguyên','mem_limit, json-file rotation, retention metrics/log và log rotation')],[4.4,11.6])
h('Firewall Windows và giới hạn mạng')
p('Kiểm tra Get-NetFirewallProfile ghi nhận cả Domain, Private và Public đang Enabled trên máy thực hành. Các trường policy mặc định trả NotConfigured nên không suy ra quy tắc hiệu lực từ chúng. Compose chỉ bind 127.0.0.1; đó là giới hạn nơi lắng nghe, không phải bằng chứng đã cấu hình UFW/iptables hoặc đã kiểm thử truy cập từ máy LAN.')
h('Ngoại lệ cần công khai')
p('cAdvisor chạy privileged và có mount host chỉ đọc để lấy số liệu Docker Desktop. Promtail chạy root với DAC_READ_SEARCH để đọc volume log, nhưng cap_drop ALL trước khi thêm quyền cụ thể. Hai ngoại lệ này phải được nêu khi demo; non-root và mount chỉ đọc chỉ là các lớp giảm rủi ro, không bảo đảm ngăn mọi leo thang đặc quyền.')
p('Không thay đổi hoặc tắt Windows Firewall trong quá trình hoàn thiện. Không mở port LAN/Internet, không tự cài CA và không công bố secret. Với triển khai thực tế cần CA tin cậy, rà soát image/quyền quan sát, chính sách sao lưu và kiểm thử tải/bảo mật riêng.')

page('CHƯƠNG 4. GIAO DIỆN VÀ LUỒNG THỰC TẾ')
h('4.1. Không gian quản trị thống nhất')
p('Khách và admin dùng chung /login với email/mật khẩu. Backend lấy vai trò từ DB để chuyển admin vào /admin hoặc khách vào /account. Route /admin/login cũ chuyển về đường dẫn chung; form đăng ký không nhận quyền admin từ client. Tài khoản bị khóa bị vô hiệu phiên khi backend tải lại trạng thái.')
fig(E/'admin-confirmation-desktop.png','Hình 4.1. Lịch demo đã xác nhận; admin quản lý đủ năm phân hệ.')
p('Trên ảnh, lịch MH-75CB5CEF4D4C thuộc Khách Demo, dịch vụ Tư vấn triển khai hệ thống, chuyên viên Phạm Gia Huy, ngày 27/10/2026 09:00–10:00. Đây là dữ liệu demo thực, được lọc theo mã để ảnh không chứa thông tin riêng của người dùng. Các thẻ tổng hợp phản ánh toàn hệ thống tại thời điểm chụp, kể cả lịch kiểm thử.')
p('Danh sách có bộ lọc mã/khách, chuyên viên, ngày và trạng thái. Nút xác nhận/hủy/hoàn tất phụ thuộc trạng thái và thời gian; server tiếp tục kiểm tra dù client tự gửi request. Hộp thoại xác nhận và thông báo lỗi giúp người dùng biết thao tác có thành công hay bị ràng buộc từ chối.')

page('4.2. KHÁCH HÀNG VÀ QUẢN LÝ DANH MỤC')
p('Khách đăng ký tài khoản, đăng nhập rồi chọn dịch vụ, chuyên viên và slot còn trống. Dữ liệu liên hệ được lấy từ hồ sơ server. Sau khi đặt, khách nhận mã lịch và xem trạng thái trong “Lịch hẹn của tôi”. Việc xác nhận được thực hiện bởi admin, không phải thông báo gửi email/SMS tự động.')
# Portrait evidence remains readable side by side; no fabricated composite UI.
t=d.add_table(rows=1,cols=2);t.autofit=False
for cell,path in zip(t.rows[0].cells,[E/'customer-account-confirmed.png',E/'admin-v3-services.png']):
    cell.width=Cm(8);cell.paragraphs[0].add_run().add_picture(str(path),width=Cm(7.2))
q=d.add_paragraph('Hình 4.2. Tài khoản khách có lịch đã xác nhận và màn hình chỉnh sửa dịch vụ trên mobile.','Caption');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Admin có thể thêm/sửa/ngừng dịch vụ; phân công nhân viên theo dịch vụ; mở lịch đơn hoặc hàng loạt, đóng/mở slot; chỉnh hồ sơ và khóa/mở tài khoản khách. Với dữ liệu đã phát sinh lịch, hệ thống bảo vệ lịch sử bằng snapshot, FK và validation. Không xóa trực tiếp bản ghi liên quan để làm “sạch” ảnh demo.')
p('Thời lượng dịch vụ là bội số 15 phút trong khoảng 15–240 phút; slot phải theo dịch vụ và không chồng lịch nhân viên. Mở hàng loạt giới hạn 32 slot cùng ngày; ngày tối đa 365 ngày. Khi chuyên viên chưa được phân công dịch vụ, admin phải phân công trước khi mở lịch.')

page('4.3. KIỂM THỬ VÀ KHẮC PHỤC LỖI')
table(['Nhóm kiểm chứng','Kết quả và nội dung'],[
('Unit','4/4: mật khẩu salted, validation, trạng thái/thời gian, token/ngày Việt Nam'),
('Integration PostgreSQL thật','3/3 nhóm: đặt trùng đồng thời, idempotency, ownership, đăng nhập và đầy đủ quản trị'),
('Stack runtime','16/16: Compose, service/job, Nginx, DB, 5 target, dashboard, LogQL'),
('Audit sâu','9/9 tại mốc đã push: GitHub, role/schema, hardening, panel có dữ liệu, nguồn log'),
('TLS','TLS 1.3, kiểm tra CA/SAN, cookie Secure, chặn metrics public'),
('Clone sạch','Đã clone GitHub, dựng 13 service với volume/cổng riêng và đạt 16 kiểm tra')],[5,11])
p('Các kiểm thử integration gửi HTTP qua Nginx và dùng DB thật. Hai request tranh cùng slot chỉ một lịch giữ được; cùng giờ ở hai chuyên viên vẫn hợp lệ. Bộ test kiểm tra khách không xem/hủy lịch của người khác, CSRF và đăng xuất; quản trị kiểm tra tạo/sửa/ngừng danh mục, phân công, mở/đóng slot và vòng đời lịch.')
h('Lỗi tìm thấy trong quá trình hoàn thiện')
p('Đăng nhập từng tách vai trò gây nhầm giữa username cũ và email khách: đã thống nhất /login và xác định vai trò ở backend. Nginx giữ IP app cũ sau recreate: đã dùng DNS resolver Docker. Sidecar xoay log bị lỗi PID namespace sau restart Nginx: đã đưa helper vào container Nginx để cùng vòng đời. Promtail thu lại file đã xoay và gán thời gian thu cho JSON: đã đọc file hiện hành, parse event time và kiểm tra mỗi request mới chỉ có một bản.')
p('Khi stack mới chạy, LogQL có thể chưa có dữ liệu dù Loki đã ready. Bộ nghiệm thu cần đợi ingestion và retry có giới hạn; không sửa kết quả failed thành passed bằng tay. Kết quả kiểm tra có thời điểm trong JSON máy cục bộ, không thay thế đánh giá demo của giảng viên.')

page('4.4. TÁI LẬP, SAO LƯU VÀ VẬN HÀNH')
h('Dựng từ checkout sạch')
p('Bản diễn tập clone trực tiếp GitHub tại commit 29cb15c, tạo .env mới với secret riêng, project Compose appointment_fresh_* và volume mới. Cổng thử là 18080/18443/15050/13001/19090 để không trùng bản chính. Toàn bộ 13 service được dựng; smoke HTTP, TLS và 16 kiểm tra stack đạt. Sau đó dừng riêng container/network bản thử; giữ volume và không đụng dữ liệu bản chính.')
p('Rollback đã diễn tập trên DB thử có một lịch fixture: image ứng viên lỗi bị readiness từ chối; quay lại image tốt, web HTTP 200 và số dòng/checksum lịch khớp. Không thử hạ schema phá tương thích trên DB chính; xem runbook release/rollback đi kèm source.')
code('node scripts/rehearse-clean-clone.js\nnode scripts/verify-stack.js\nnode scripts/verify-tls.js\nnode scripts/audit-requirements.js')
h('Backup và kiểm tra phục hồi')
p('Đã tạo bản dump appointment-20261024-051045.dump và restore vào DB riêng restore_verify_20261024051048. Schema 1/2/3 và 26 lịch tại thời điểm backup được giữ; số dòng và checksum của 8 bảng nghiệp vụ khớp. Đây là kiểm tra một bản sao tại thời điểm đó, không phải cơ chế sao lưu tự động hằng ngày.')
code('powershell -ExecutionPolicy Bypass -File scripts/backup.ps1\npowershell -ExecutionPolicy Bypass -File scripts/restore-test.ps1 `\n  -BackupFile backups/appointment-YYYYMMDD-HHMMSS.dump')
p('SHA256 bản dump được lưu trong evidence cục bộ để đối chiếu tính toàn vẹn. File backup bị gitignore vì có thể chứa thông tin tài khoản; không đính kèm báo cáo. Thử restore vào DB tách biệt, không ghi đè DB chính. App recreate đã kiểm chứng giữ dữ liệu và phiên đăng nhập nhờ PostgreSQL/volume.')
h('Xử lý khi không truy cập được')
p('Kiểm tra Docker Engine, docker compose ps -a và healthcheck trước. Nếu 8080 lỗi, xem log Nginx/app; nếu login sai, đối chiếu tài khoản DB với .env thay vì xóa volume. Nếu dashboard trống, kiểm tra targets và đợi đủ scrape; nếu log thiếu, kiểm tra Promtail/Loki và thời gian query. Không dùng down -v như cách sửa lỗi mặc định.')

page('CHƯƠNG 5. ĐÁNH GIÁ VÀ HƯỚNG PHÁT TRIỂN')
h('5.1. Kết quả đạt được')
p('Hệ thống có đủ bốn đối tượng chính của đề: dịch vụ, khung giờ, khách hàng và xác nhận, cùng chuyên viên và phân công phục vụ mô hình nhiều nhân viên. Frontend và backend dùng chung PostgreSQL thật; admin và khách có quyền riêng. Hạ tầng Compose, Nginx, pgAdmin, monitoring, logging, hardening, GitHub và hồ sơ báo cáo đã được tổ chức để chạy và demo lại trên localhost.')
h('5.2. Giới hạn')
p('Chưa có tài khoản riêng cho nhân viên, xác minh email, quên mật khẩu, email/SMS, thanh toán hay đồng bộ lịch ngoài. Danh sách có giới hạn số dòng và bộ lọc, chưa phân trang đầy đủ. Log tập trung mới thu app/Nginx/PostgreSQL; lịch sử Loki có thể còn dữ liệu trùng trước mốc sửa. Chưa kiểm thử tải lớn, HA hoặc phục hồi thảm họa trên máy khác. Windows Firewall được kiểm tra trạng thái bật, chưa nghiệm thu quy tắc từ máy LAN; không có UFW trên host Windows.')
p('HTTPS dùng chứng chỉ tự ký, không phải CA công cộng. cAdvisor privileged là ngoại lệ quan sát được công khai. Tên repository chứa MSSV, tài khoản GitHub vẫn là vzmanhdev theo người dùng cung cấp; không kết luận username tài khoản đã trùng MSSV. Điểm số và chấp nhận bài thuộc giảng viên dựa trên yêu cầu và demo.')
h('5.3. Hướng mở rộng sau bài thực hành')
p('Ưu tiên tài khoản nhân viên chỉ xem lịch được giao, email xác minh và nhắc lịch; thêm phân trang và chính sách lịch nghỉ. Với Internet cần CA tin cậy, rà soát image/quyền quan sát, thay pipeline log theo vòng đời công cụ, kế hoạch backup tự động và kiểm thử khôi phục định kỳ. Những nội dung này không được ghi là đã triển khai trong phạm vi localhost hiện tại.')
h('Tài liệu tham khảo')
p('[1] Đề tài Thực hành – Triển khai và Quản trị Hệ thống Phần mềm (Đề 40), tài liệu học phần do người dùng cung cấp. [2] Chapter 01–09, học phần; dùng làm tài liệu kiến thức. [3] Bao_Cao_Trien_Khai_He_Thong.docx: khung đối chiếu. [4] Bao_cao_Appointment_System.docx: chỉ tham chiếu bố cục bìa. [5] Source, cấu hình và README của repository GitHub dự án; kết quả kiểm thử ngày 24/10/2026.')

page('PHỤ LỤC. KỊCH BẢN DEMO VÀ BÀN GIAO')
table(['Thời gian','Thao tác và kết quả mong đợi'],[
('0–1 phút','Giới thiệu đề 40, mô hình khách–chuyên viên, docker compose ps -a'),
('1–4 phút','Admin tạo dịch vụ/nhân viên, phân công và mở slot tương lai'),
('4–6 phút','Khách đăng ký/đăng nhập, chọn dịch vụ–người–giờ, đặt pending'),
('6–8 phút','Admin lọc mã, xác nhận; khách tải lại thấy confirmed; thử slot trùng bị từ chối'),
('8–9 phút','pgAdmin xem services, staff, time_slots, appointments, events'),
('9–11 phút','Prometheus 5 target; Grafana CPU/RAM/I/O và HTTP/DB; giải thích dữ liệu'),
('11–13 phút','Explore chạy ba LogQL; tạo request 404 và xem log mới'),
('13–15 phút','Chạy kiểm tra UID/network/TLS; trình bày Git, backup và giới hạn')],[3.1,12.9])
h('Bộ lệnh kiểm chứng nhanh')
code('docker compose ps -a\ndocker compose exec -T nginx nginx -t\ndocker compose exec -T app id\nnode scripts/verify-stack.js\nnode scripts/verify-tls.js\nnpm.cmd test\nnpm.cmd run test:integration\nnode scripts/audit-requirements.js\ngit log --oneline -12\ngit status --short')
p('Các URL chính: http://localhost:8080; https://localhost:8443; pgAdmin :5050; Grafana :3001/d/appointment-local; Prometheus :9090/targets. Tài khoản admin và công cụ lấy từ .env trên máy; đăng ký tài khoản khách để demo. Dữ liệu kiểm thử dùng địa chỉ example.com, không gửi thư thật. Có thể quay màn hình theo kịch bản này khi cần nộp video; báo cáo không nhận là đã có video ghi hình nếu chưa thực hiện.')
p('Bàn giao gồm source/config trên GitHub, README, hướng dẫn quản trị v3, ma trận đối chiếu, kịch bản demo và báo cáo Word cá nhân. Trước khi nộp, người thực hiện xem lại thông tin giảng viên ở bìa và yêu cầu riêng của lớp; không đưa .env, private key hoặc dump DB vào bài nộp công khai.')

# Remove unused sample images/hyperlinks while retaining the actual cover logo.
used=set()
for part in [d.part,*[s.footer.part for s in d.sections],*[s.header.part for s in d.sections]]:
    for el in part.element.iter():
        for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
            if el.get(attr):used.add(el.get(attr))
for rid,rel in list(d.part.rels.items()):
    if rid not in used and rel.reltype.rsplit('/',1)[-1] in ['image','hyperlink']:d.part.drop_rel(rid)
settings=d.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
d.save(OUT)
assert sha256(REF.read_bytes()).hexdigest()==fingerprint
(TMP/'build-manifest.json').write_text(json.dumps({'reference_sha256':fingerprint,'output':str(OUT),'planned_pages':len(pages)+1,'titles':pages},ensure_ascii=False,indent=2),encoding='utf8')
print('Created',OUT,'planned pages',len(pages)+1)
